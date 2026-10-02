from io import BytesIO
from pathlib import Path
import math

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "QuickAssist_SDS_TheoMau_TrinhBayGiongMau.docx"
NAVY, BLUE, LIGHT_BLUE, GRAY, WHITE = "17365D", "2F75B5", "D9EAF7", "F2F2F2", "FFFFFF"


def font(size, bold=False):
    try:
        return ImageFont.truetype("arialbd.ttf" if bold else "arial.ttf", size)
    except OSError:
        return ImageFont.load_default()


def center_text(draw, box, text, size=22, color="#17365D", bold=True):
    f = font(size, bold)
    b = draw.multiline_textbbox((0, 0), text, font=f, align="center", spacing=4)
    draw.multiline_text(((box[0]+box[2]-(b[2]-b[0]))/2, (box[1]+box[3]-(b[3]-b[1]))/2), text, font=f, fill=color, align="center", spacing=4)


def box(draw, x, y, w, h, text, fill="#D9EAF7", size=22):
    bounds = (x, y, x+w, y+h)
    draw.rounded_rectangle(bounds, radius=18, fill=fill, outline="#17365D", width=4)
    center_text(draw, bounds, text, size=size)


def arrow(draw, start, end, color="#2F75B5", width=4):
    draw.line([start, end], fill=color, width=width)
    a = math.atan2(end[1]-start[1], end[0]-start[0])
    n = 16
    draw.polygon([end, (end[0]-n*math.cos(a-.5), end[1]-n*math.sin(a-.5)), (end[0]-n*math.cos(a+.5), end[1]-n*math.sin(a+.5))], fill=color)


def double_arrow(draw, start, end, color="#E3C400", width=8):
    """Two-way yellow arrow, matching the supplied SDS architecture style."""
    draw.line([start, end], fill="#B89200", width=width + 3)
    draw.line([start, end], fill=color, width=width)
    for tip, tail in ((start, end), (end, start)):
        a = math.atan2(tip[1]-tail[1], tip[0]-tail[0]); n = 26
        points = [tip, (tip[0]-n*math.cos(a-.52), tip[1]-n*math.sin(a-.52)), (tip[0]-n*math.cos(a+.52), tip[1]-n*math.sin(a+.52))]
        draw.polygon(points, fill="#B89200")
        inner = [tip, (tip[0]-(n-4)*math.cos(a-.52), tip[1]-(n-4)*math.sin(a-.52)), (tip[0]-(n-4)*math.cos(a+.52), tip[1]-(n-4)*math.sin(a+.52))]
        draw.polygon(inner, fill=color)


def dashed_rectangle(draw, bounds, color="#555555", width=3, dash=12):
    x1, y1, x2, y2 = bounds
    for horizontal, fixed, a, b in [(True,y1,x1,x2),(True,y2,x1,x2),(False,x1,y1,y2),(False,x2,y1,y2)]:
        p = a
        while p < b:
            e = min(p+dash,b)
            draw.line((p,fixed,e,fixed) if horizontal else (fixed,p,fixed,e), fill=color, width=width)
            p += dash*2


def cylinder(draw, x, y, w, h, label):
    draw.rectangle((x,y+h*.16,x+w,y+h*.84),fill="#FBD9AA",outline="#B97920",width=3)
    draw.ellipse((x,y,x+w,y+h*.32),fill="#FCE3C2",outline="#B97920",width=3)
    draw.arc((x,y+h*.68,x+w,y+h),0,180,fill="#B97920",width=3)
    center_text(draw,(x,y,x+w,y+h),label,18,"#1A1A1A",False)


def png_image(draw_fn, width=1800, height=780):
    img = Image.new("RGB", (width, height), "white")
    draw_fn(ImageDraw.Draw(img), width, height)
    bio = BytesIO(); img.save(bio, format="PNG"); bio.seek(0)
    return bio


def logical_architecture():
    def paint(d, w, h):
        left, right, layer_w, layer_h = 70, 1320, 1250, 145
        layers = [(110,"Presentation/Integration Layer","#A9D9EF"),(300,"Business Layer","#F7C9CC"),(490,"Data Access Layer","#FFFF92"),(680,"Data Layer","#C9E994")]
        for y,title,color in layers:
            d.rectangle((left,y,right,y+layer_h),fill=color,outline="#4E4E4E",width=3)
            d.text((left+8,y+12),title,font=font(20,True),fill="#1A1A1A")
        # Infrastructure column on the right.
        d.rectangle((1440,110,1710,715),fill="#F9D67B",outline="#C59A36",width=3)
        # Pillow cannot rotate text inline; use stacked words for the vertical side layer label.
        d.text((1510,180),"Infrastructure",font=font(20,True),fill="#1A1A1A")
        d.text((1510,210),"Layer",font=font(20,True),fill="#1A1A1A")
        for x,y,tw,th,label in [
            (150,180,250,65,"ExtensionUI"),(535,180,250,65,"GoogleOAuthService"),(920,180,250,65,"NotesService"),
            (305,370,275,65,"SearchSkimService"),(780,370,250,65,"AIRAGService"),
            (565,560,280,65,"RepositoryService"),(610,730,230,60,"PostgreSQL + pgvector"),(1490,395,170,65,"LoggingService")]:
            box(d,x,y,tw,th,label,"#FBD3A0",18)
        # Layer interchanges and infrastructure exchanges.
        double_arrow(d,(695,255),(695,300)); double_arrow(d,(695,445),(695,490)); double_arrow(d,(695,635),(695,680))
        for y in (215,405,595): double_arrow(d,(1320,y),(1440,y))
    return png_image(paint, 1800, 900)


def physical_architecture():
    def paint(d, w, h):
        dashed_rectangle(d,(85,70,690,715),"#555555",3,14); d.text((387,90),"External Service",font=font(23,False),fill="#1A1A1A",anchor="mm")
        for y,label in [(150,"Browser Clients"),(330,"Google OAuth Servers"),(510,"AI Provider Servers")]:
            # Server icon approximation in the same dark, photographic placement.
            d.rounded_rectangle((300,y,450,y+105),radius=8,fill="#253341",outline="#111111",width=3)
            d.rectangle((325,y+22,425,y+37),fill="#5E7181"); d.rectangle((325,y+52,425,y+67),fill="#5E7181")
            d.ellipse((334,y+78,347,y+91),fill="#66A3D2"); d.ellipse((357,y+78,370,y+91),fill="#66A3D2")
            d.text((375,y-25),label,font=font(20,False),fill="#1A1A1A",anchor="mm")
        d.rectangle((820,120,1710,715),outline="#333333",width=3); d.text((855,155),"Application Server 1",font=font(20,False),fill="#1A1A1A")
        d.text((855,350),"Application Server 2",font=font(20,False),fill="#1A1A1A")
        d.text((855,540),"Application Server n",font=font(20,False),fill="#1A1A1A")
        for y in (205,400,590):
            d.rounded_rectangle((1030,y,1170,y+85),radius=7,fill="#253341",outline="#111111",width=3)
            d.rectangle((1050,y+20,1150,y+33),fill="#5E7181"); d.rectangle((1050,y+48,1150,y+60),fill="#5E7181")
            arrow(d,(1170,y+42),(1450,420),"#202020",3)
        d.rectangle((1450,365,1515,470),fill="#2D82E6",outline="#9AD1FF",width=3); d.polygon([(1450,365),(1483,337),(1515,365),(1483,392)],fill="#77B8FF")
        d.text((1482,325),"PostgreSQL database",font=font(19,False),fill="#1A1A1A",anchor="mm")
        arrow(d,(690,395),(820,395),"#202020",3)
    return png_image(paint)


def component_architecture():
    def paint(d, w, h):
        # Central pale-yellow project container, blue internal services, orange external services.
        d.rectangle((300,120,1360,655),fill="#FFF1C7",outline="#4E4E4E",width=3); d.text((320,140),"QuickAssist Project",font=font(25,True),fill="#1A1A1A")
        d.rectangle((70,190,240,640),fill="#8E8E8E",outline="#666666",width=2); d.text((155,210),"Data storage",font=font(17,True),fill="#1A1A1A",anchor="mm")
        box(d,95,360,120,120,"PostgreSQL\npgvector","#D7D7D7",17)
        # External top bars and right nodes.
        d.rectangle((300,55,800,95),fill="#F5B27D",outline="#B77946",width=2); d.text((550,75),"Google OAuth / OIDC",font=font(20,False),fill="#1A1A1A",anchor="mm")
        d.rectangle((800,55,1360,95),fill="#F5B27D",outline="#B77946",width=2); d.text((1080,75),"AI Provider (Embedding / Rerank)",font=font(20,False),fill="#1A1A1A",anchor="mm")
        services=[(390,325,220,150,"Notes\nService"),(705,210,230,150,"GoogleOAuth\nService"),(1010,210,230,150,"Extension\nGateway"),(1010,445,230,150,"SearchSkim\nService"),(705,445,230,150,"AIRAG\nService")]
        for x,y,sw,sh,label in services: d.rectangle((x,y,x+sw,y+sh),fill="#18A7E0",outline="#1683B0",width=2); center_text(d,(x,y,x+sw,y+sh),label,21,"#1A1A1A",False)
        for y,label in [(245,"Google OAuth"),(490,"AI Provider")]: box(d,1470,y,180,90,label,"#F5B27D",19)
        for s,e in [((240,420),(390,400)),((610,400),(705,520)),((610,400),(705,285)),((935,285),(1010,285)),((935,520),(1010,520)),((1240,285),(1470,290)),((1240,520),(1470,535)),((815,210),(815,95)),((820,445),(1060,95))]: arrow(d,s,e,"#2D5AA5",3)
    return png_image(paint)


def sequence_diagram():
    def paint(d, w, h):
        actors=[(90,"user"),(310,"Extension"),(560,"OAuth\nService"),(810,"Notes\nService"),(1080,"AI\nProvider"),(1340,"Database")]
        for x,label in actors:
            d.rectangle((x-68,75,x+68,115),fill="white",outline="#222222",width=2); center_text(d,(x-68,75,x+68,115),label,13,"#111111",False)
            d.line((x,115,x,730),fill="#555555",width=2)
        events=[(155,90,310,"click Sign in / ghi chú"),(205,310,560,"OAuth code + PKCE"),(255,560,310,"app session"),(310,310,810,"POST /notes/quick"),(365,810,1340,"save note"),(420,810,1080,"chunk / embed"),(475,1080,810,"vectors"),(530,810,1340,"save chunks + quota"),(585,810,310,"202 PROCESSING"),(650,310,90,"status notification")]
        for y,x1,x2,label in events:
            arrow(d,(x1,y),(x2,y),"#1A1A1A",2); d.text(((x1+x2)//2,y-13),label,font=font(13,False),fill="#111111",anchor="mm")
        # UML-style alt frame for quota/provider error branch.
        d.rectangle((730,555,1230,685),outline="#222222",width=2); d.rectangle((730,555,790,580),fill="white",outline="#222222",width=2); d.text((760,568),"alt",font=font(12,False),fill="#111111",anchor="mm")
        d.text((800,605),"[quota đủ]  → READY",font=font(13,False),fill="#111111")
        d.line((730,620,1230,620),fill="#555555",width=1)
        d.text((800,650),"[quota thiếu / provider lỗi]  → SKIPPED hoặc FAILED",font=font(13,False),fill="#111111")
    return png_image(paint, 1700, 790)


def shade(cell, color):
    p = cell._tc.get_or_add_tcPr(); e = OxmlElement("w:shd"); e.set(qn("w:fill"), color); p.append(e)


def cell_text(cell, value, bold=False, color=None, size=9):
    cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(value)); r.font.name = "Times New Roman"; r.font.size = Pt(size); r.bold = bold
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, size=8.5):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell_text(t.rows[0].cells[i],h,True,"1A1A1A",9); shade(t.rows[0].cells[i],"E7E6E6")
    for n,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cell_text(cells[i],v,size=size)
            if n%2: shade(cells[i],GRAY)
    doc.add_paragraph(); return t


def heading(doc, text, level=1):
    p=doc.add_heading(text,level); p.paragraph_format.keep_with_next=True; return p


def bullets(doc, items):
    for item in items: doc.add_paragraph(item,style="List Bullet")


def field(run, code):
    a=OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"),"begin")
    b=OxmlElement("w:instrText"); b.set(qn("xml:space"),"preserve"); b.text=code
    c=OxmlElement("w:fldChar"); c.set(qn("w:fldCharType"),"end")
    run._r.extend([a,b,c])


doc=Document(); sec=doc.sections[0]
sec.top_margin=Cm(2); sec.bottom_margin=Cm(1.7); sec.left_margin=Cm(2.1); sec.right_margin=Cm(2.1)
styles=doc.styles
styles["Normal"].font.name="Times New Roman"; styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"),"Times New Roman"); styles["Normal"].font.size=Pt(11); styles["Normal"].paragraph_format.line_spacing=1.15
for name,size,color in [("Title",24,"000000"),("Heading 1",16,"000000"),("Heading 2",13,"000000"),("Heading 3",11,"000000")]:
    st=styles[name]; st.font.name="Times New Roman"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"Times New Roman"); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)

# The reference PDF uses a clean academic page with only a centred page number.
sec.header.paragraphs[0].text=""
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; field(fp.add_run(),"PAGE")

# Cover follows the reference PDF.
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(105)
r=p.add_run("Tài liệu Đặc tả Thiết kế Giải pháp\n(Solution Design Specification - SDS)"); r.font.size=Pt(23); r.font.bold=True; r.font.color.rgb=RGBColor.from_string("000000")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(25)
r=p.add_run("QuickAssist – Trợ lý Ghi chú nhanh dựa trên AI"); r.font.size=Pt(19); r.font.bold=True; r.font.color.rgb=RGBColor.from_string("000000")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(45)
p.add_run("Người A, Người B, Người C, Người D\n04/09/2026").font.size=Pt(13)
doc.add_page_break()

heading(doc,"Mục lục",1); p=doc.add_paragraph(); field(p.add_run(),'TOC \\o "1-3" \\h \\z \\u'); doc.add_paragraph("Khi mở bằng Microsoft Word, chọn Update Field/Update Table để cập nhật số trang.").runs[0].italic=True; doc.add_page_break()

heading(doc,"Quản lý Tài liệu",1)
table(doc,["Trường","Nội dung"],[
    ("Phiên bản","1.0"),("Trạng thái","Dự thảo (Draft)"),("Ngày","04/09/2026"),("Tác giả","Người A, Người B, Người C, Người D"),("Giảng viên hướng dẫn","TS. Nguyễn Trọng Phúc"),("Mã tài liệu","QA-SDS-002"),
])
doc.add_paragraph("Bảng 1: Lịch sử Phiên bản",style=None)
table(doc,["Phiên bản","Ngày","Tác giả","Nội dung Thay đổi"],[
    ("1.0","04/09/2026","Nhóm 1","Viết lại theo mẫu SDS; xác thực dùng Google OAuth; loại bỏ lịch trình cá nhân."),
])
doc.add_page_break()

heading(doc,"1 Giới thiệu",1)
heading(doc,"1.1 Mục đích",2); doc.add_paragraph("Tài liệu cung cấp thiết kế kỹ thuật chi tiết cho QuickAssist, một phần mở rộng trình duyệt giúp người dùng ghi chú nhanh nội dung web, quản lý thư viện ghi chú và tra cứu theo ngữ nghĩa. Tài liệu mô tả kiến trúc, thành phần, luồng xử lý, dữ liệu và công nghệ để nhóm triển khai thống nhất.")
heading(doc,"1.2 Tổng quan Dự án",2); doc.add_paragraph("QuickAssist cho phép người dùng bôi đen văn bản trên trang web để lưu ghi chú, duyệt/sửa/xóa note đã lưu, tìm kiếm trong thư viện bằng hybrid search và tìm các đoạn nổi bật trong trang. Extension sử dụng Google OAuth để xác thực; server điều phối lưu trữ, quota và dịch vụ AI.")
heading(doc,"1.3 Phạm vi",2); heading(doc,"1.3.1 Trong Phạm vi",3); bullets(doc,["Đăng nhập bằng Google OAuth 2.0 Authorization Code + PKCE với scope openid, email, profile.","Ghi chú nhanh từ selection, có popup đếm ngược và hoàn tác.","Quản lý note, phân mục, CRUD, cache IndexedDB và đồng bộ server.","Tìm kiếm RAG-like theo user/category bằng BM25, vector similarity và reranking.","Tìm đoạn nổi bật trong trang bằng Readability, session TTL và lưu các đoạn được chọn.","Quota AI, logging, retry, cấu hình AI Provider và triển khai prototype."])
heading(doc,"1.3.2 Ngoài Phạm vi",3); bullets(doc,["Trích xuất, xác nhận hoặc quản lý lịch trình cá nhân.","Đăng nhập bằng mật khẩu, đăng ký tài khoản thủ công hoặc lưu mật khẩu Google.","Chia sẻ note đa người dùng, CRM bên thứ ba và fine-tuning mô hình AI.","Tự động gửi hành động thay người dùng trên nội dung web."])
heading(doc,"1.4 Tài liệu Tham khảo",2); bullets(doc,["Báo cáo ý tưởng đề tài QuickAssist.","Kịch bản chức năng QuickAssist.","Kế hoạch phân công QuickAssist — bản ẩn danh Google OAuth.","Google OAuth 2.0 / OpenID Connect documentation; PostgreSQL/pgvector documentation."])

heading(doc,"2 Giả định và Ràng buộc",1)
heading(doc,"2.1 Giả định",2); bullets(doc,["Người dùng có tài khoản Google hợp lệ và chấp thuận scope tối thiểu.","Google OAuth client ID, redirect URI theo môi trường và AI Provider đã được cấu hình an toàn.","Trình duyệt mục tiêu là Chrome/Edge hỗ trợ Manifest V3.","PostgreSQL có extension pgvector; server có thể truy cập AI Provider qua private network hoặc HTTPS."])
heading(doc,"2.2 Ràng buộc",2); bullets(doc,["Công nghệ cốt lõi: TypeScript/Manifest V3, Python API, PostgreSQL/pgvector và AI adapter.","Google OAuth, AI Provider và trình duyệt đều chịu rate limit/quota.","Extension không chứa client secret hay provider secret; server là nơi đổi OAuth code và phát session.","Prototype triển khai monolith để phù hợp phạm vi môn học, nhưng các module phải tách rời về contract."])

heading(doc,"3 Yêu cầu",1)
heading(doc,"3.1 Yêu cầu Chức năng (Functional Requirements)",2)
table(doc,["Mã","Yêu cầu"],[
    ("FR-01","Đăng nhập qua Google OAuth với PKCE; server xác minh state, nonce, issuer, audience, expiry và phát app session."),
    ("FR-02","Người dùng có thể bôi đen văn bản, mở Ghi chú nhanh, hoàn tác trong 5 giây và lưu note có URL/title."),
    ("FR-03","Hệ thống lưu note theo user; AI chunk/embed nếu đủ quota; hiển thị trạng thái SAVED/PROCESSING/READY/SKIPPED/FAILED."),
    ("FR-04","Người dùng duyệt, lọc, sửa, xóa một hoặc nhiều note; extension cache-first bằng IndexedDB và đồng bộ incremental."),
    ("FR-05","Người dùng tìm note theo toàn thư viện hoặc phân mục; server chỉ search pool thuộc user đó và trả top-k chunk."),
    ("FR-06","Người dùng tìm nội dung trong trang qua Readability; session có TTL và chỉ lưu lâu dài các đoạn user chọn."),
    ("FR-07","Hệ thống ghi quota ledger, requestId và log kỹ thuật không chứa token/nội dung nhạy cảm."),
])
heading(doc,"3.2 Yêu cầu Phi chức năng (Non-functional Requirements)",2)
bullets(doc,["Performance: popup phản hồi < 200 ms; CRUD API p95 < 500 ms; RAG p95 mục tiêu < 8 giây ở cấu hình demo.","Security: mọi API dữ liệu xác thực bằng app access token và kiểm tra ownership; không lưu mật khẩu Google.","Reliability: retry theo exponential backoff cho lỗi mạng/5xx; idempotency key tránh tạo note trùng.","Maintainability: API versioned /api/v1, OpenAPI, biến cấu hình bên ngoài và migration versioned.","Usability: có trạng thái loading/empty/error/undo rõ ràng; thao tác chính hỗ trợ keyboard.","Resource constraints: giới hạn payload, timeout, concurrency và quota cho Google OAuth/AI/database."])
doc.add_page_break()

heading(doc,"4 Kiến trúc Giải pháp",1)
heading(doc,"4.1 Sơ đồ kiến trúc logic",2); doc.add_picture(logical_architecture(),width=Inches(6.8)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER; doc.add_paragraph("Hình 1: Sơ đồ kiến trúc logic",style=None).alignment=WD_ALIGN_PARAGRAPH.CENTER
heading(doc,"4.2 Sơ đồ kiến trúc vật lý",2); doc.add_paragraph("Với phạm vi môn học, hệ thống triển khai dạng monolith trên một máy chủ hoặc Docker Compose. Extension giao tiếp qua HTTPS; database và AI adapter đặt sau API. Khi cần mở rộng, API/worker có thể tách thành nhiều instance, còn database dùng connection pool và backup riêng.")
doc.add_picture(physical_architecture(),width=Inches(6.8)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER; doc.add_paragraph("Hình 2: Sơ đồ kiến trúc vật lý",style=None).alignment=WD_ALIGN_PARAGRAPH.CENTER
heading(doc,"4.3 Sơ đồ Kiến trúc các thành phần",2); doc.add_picture(component_architecture(),width=Inches(6.8)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER; doc.add_paragraph("Hình 3: Thành phần kiến trúc tổng thể",style=None).alignment=WD_ALIGN_PARAGRAPH.CENTER
bullets(doc,["Hệ thống bên ngoài: Google OAuth Authorization Server và AI Provider.","Ứng dụng lõi: Browser Extension, Google OAuth Module, Notes & Sync Module, Search & Skim Module, AI Adapter.","Kho dữ liệu: PostgreSQL + pgvector, IndexedDB trên client, TTL store cho skim session.","Tương tác chính: User → Extension → API → AI/Data; server phát app session sau Google OAuth."])

heading(doc,"5 Thiết kế Chi tiết",1)
heading(doc,"5.1 Sequence Diagram",2); doc.add_picture(sequence_diagram(),width=Inches(6.8)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER; doc.add_paragraph("Hình 4: Luồng ghi chú nhanh và xử lý RAG",style=None).alignment=WD_ALIGN_PARAGRAPH.CENTER
bullets(doc,["User bôi đen nội dung, extension tạo popup undo 5 giây và pending action có idempotency key.","API xác thực app session, kiểm tra ownership/category rồi lưu note; việc AI xử lý không giữ transaction DB.","AI trả chunk/vector; server lưu vector, cập nhật quota ledger và đồng bộ trạng thái để client hiển thị.","Nếu quota không đủ, note vẫn được lưu với trạng thái SKIPPED; nếu lỗi provider, trạng thái FAILED và có retry kiểm soát."])
heading(doc,"5.2 GoogleOAuthService (Xác thực và Phân quyền)",2); bullets(doc,["Công nghệ: Google OAuth 2.0/OpenID Connect, browser identity flow, Python JWT/OIDC validation.","Kích hoạt: user chọn đăng nhập trên extension; extension dùng Authorization Code Flow + PKCE.","Chức năng: tạo/kiểm tra state và nonce, đổi code, xác minh ID token, ánh xạ google_sub sang users, phát app access/refresh token, logout local session.","Giao diện tương tác: Google Authorization Server, Extension API client, session store/users table."])
heading(doc,"5.3 NotesSyncService (Ghi chú, CRUD và đồng bộ)",2); bullets(doc,["Công nghệ: Python ASGI API, PostgreSQL, IndexedDB API ở extension.","Chức năng: tạo note nhanh, GET/PUT/DELETE/bulk-delete, ownership check, version conflict, cache-first sync và pending action retry.","Giao diện tương tác: Extension API client, Notes/Category repository, QuotaService và AIAdapter."])
heading(doc,"5.4 AIRAGService (Chunk, Embedding, Retrieval và Rerank)",2); bullets(doc,["Công nghệ: AI adapter tới Dify/local provider, PostgreSQL pgvector, BM25 implementation.","Chức năng: semantic chunking, embedding, HNSW/IVFFlat vector search, BM25, kết hợp kết quả, rerank/fallback và đo latency.","Giao diện tương tác: NotesSyncService, SearchSkimService, AI Provider và note_chunks repository."])
heading(doc,"5.5 SearchSkimService (Tìm kiếm và nội dung trang)",2); bullets(doc,["Công nghệ: Mozilla Readability trong extension, API session TTL và hybrid retrieval.","Chức năng: làm sạch text trang, tạo skim session theo user/tab, search trong session, trả preview và chuyển selected chunks thành note.","Dữ liệu skim chưa được user chọn chỉ tồn tại trong TTL store; session hết hạn trả HTTP 410."])

heading(doc,"6 Thiết kế Dữ liệu",1)
table(doc,["Bảng","Trường chính","Mục đích"],[
    ("users","id UUID PK, google_sub UNIQUE, email UNIQUE, email_verified, display_name, avatar_url, status","Liên kết tài khoản Google; không lưu mật khẩu"),
    ("user_preferences","user_id FK, auto_process, notification_flags, updated_at","Tùy chọn của người dùng"),
    ("categories","id UUID PK, user_id FK, name, color, version","Phân mục note thuộc user"),
    ("notes","id UUID PK, user_id FK, category_id, source_url/title, content, status, version, timestamps","Note chuẩn, soft delete để đồng bộ"),
    ("note_chunks","id UUID PK, note_id/user_id FK, ordinal, text, embedding VECTOR, model_version, metadata","Chunk/vector cho RAG"),
    ("quota_accounts / quota_ledger","user_id, remaining_units / operation, units, request_id","Quota hiện hành và audit append-only"),
    ("idempotency_keys","user_id, key, request_hash, expires_at","Chống tạo trùng do retry"),
    ("skim_sessions","session_id, user_id, tab_ref_hash, expires_at","Metadata session; text ở TTL store"),
])
bullets(doc,["Index chính: notes(user_id, updated_at), notes(user_id, category_id), note_chunks(user_id, note_id) và vector index theo benchmark.","Mọi query note/chunk/category có điều kiện user_id; user_id lấy từ auth context, không tin từ JSON client.","Khi note đổi nội dung, vector cũ chỉ thay sau khi chunks mới được tạo thành công."])

heading(doc,"7 Công nghệ Sử dụng",1)
bullets(doc,["Browser Extension: TypeScript, Manifest V3, WebExtension APIs, IndexedDB và Mozilla Readability.","Xác thực: Google OAuth 2.0/OpenID Connect, Authorization Code + PKCE.","Backend: Python + ASGI framework, OpenAPI, migration framework và worker nền.","Dữ liệu: PostgreSQL, pgvector; Redis-compatible TTL store tùy môi trường.","AI: Dify hoặc llama.cpp/vLLM qua provider adapter; mock provider cho test.","Triển khai: Docker Compose, biến môi trường/secret store, Git-based CI kiểm tra build/lint/test."])

heading(doc,"8 Bảo mật",1)
bullets(doc,["Google OAuth: scope tối thiểu openid/email/profile; validate state, nonce, PKCE verifier, issuer, audience và expiry. Không nhúng client secret trong extension.","Session: app access token ngắn hạn, refresh token rotation/revocation; logout chỉ revoke local session QuickAssist.","Authorization: mọi API có auth context và ownership check; 404 được dùng để che tài nguyên không thuộc user.","Input: giới hạn body/ký tự, sanitize text/HTML, không render innerHTML không kiểm soát; page text là untrusted input với AI.","Secrets: DATABASE_URL, OAuth config và AI key ở environment/secret store; không hard-code và không ghi log.","Privacy: chỉ thu thập selection/page text sau thao tác rõ ràng; skim text có TTL, log không chứa raw note/token."])

heading(doc,"9 Triển khai",1)
bullets(doc,["Môi trường prototype: Docker Compose gồm API/worker, PostgreSQL+pgvector và AI adapter/provider tùy chọn.","Cấu hình chính: GOOGLE_CLIENT_ID, OAUTH_REDIRECT_URI, GOOGLE_ISSUERS, APP_TOKEN_TTL, DATABASE_URL, AI_PROVIDER, AI_BASE_URL, SESSION_TTL, MAX_PAGE_CHARS.","Quy trình: migration DB → khởi động API/AI → health check → build extension với API base URL → smoke test Google OAuth, quick note, CRUD, RAG và skim.","Khả năng mở rộng: tách worker AI, chạy nhiều API instance, connection pool DB, rate limit/circuit breaker cho provider."])

heading(doc,"10 Xử lý Lỗi và Giảm thiểu Rủi ro",1)
table(doc,["Rủi ro / lỗi","Biện pháp Giảm thiểu"],[
    ("Google OAuth state/nonce/token không hợp lệ","Reject request, log requestId an toàn, yêu cầu đăng nhập lại; test issuer/audience/expiry và redirect URI theo môi trường."),
    ("AI provider timeout hoặc lỗi 5xx","Timeout, exponential backoff có jitter, circuit breaker; rerank fallback về hybrid score; note giữ FAILED để retry."),
    ("Quota không đủ","Ghi ledger nguyên tử; note vẫn SAVED nhưng SKIPPED; UI nêu rõ lý do và không trừ quota khi silent skip."),
    ("Truy cập chéo note user","Repository bắt buộc user_id; IDOR integration test; trả 404 thay vì lộ tồn tại tài nguyên."),
    ("Đồng bộ IndexedDB/server xung đột","version/updatedAt, 409 conflict, optimistic update và option tải bản mới/lưu bản sao."),
    ("Readability/trang dài không ổn định","Giới hạn payload, bộ trang test, fallback và thông báo khi không trích được nội dung."),
    ("Rate limit/DB không khả dụng","429/5xx retry có giới hạn, connection pool, health/readiness checks và log/metric theo stage."),
])

doc.core_properties.title="QuickAssist — SDS theo mẫu Solution Design Specification"; doc.core_properties.author="Nhóm 1"; doc.core_properties.subject="Google OAuth, ghi chú nhanh và RAG"; doc.core_properties.keywords="QuickAssist,SDS,Google OAuth,Manifest V3,RAG,pgvector"
doc.save(OUTPUT)
check=Document(OUTPUT); assert len(check.tables)>=5 and len(check.inline_shapes)==4
print(f"Created: {OUTPUT}"); print(f"Tables: {len(check.tables)}; figures: {len(check.inline_shapes)}; paragraphs: {len(check.paragraphs)}")
