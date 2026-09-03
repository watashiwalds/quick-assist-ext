from pathlib import Path
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "QuickAssist_KeHoach_PhanCong_4Nguoi.xlsx"
OUTPUT = ROOT / "QuickAssist_KeHoach_PhanCong_AnDanh_A_B_C_D.xlsx"

# Only identities are replaced. Roles, task ownership, collaboration, estimates,
# timeline, risks, formulas, and all allocation data remain unchanged.
mapping = {
    "Mai Đức Vinh": "Người A",
    "Trịnh Mạnh Quang": "Người B",
    "Lê Đăng Sơn": "Người C",
    "Trần Thuỳ Dương": "Người D",
    "Vinh": "Người A",
    "Quang": "Người B",
    "Sơn": "Người C",
    "Dương": "Người D",
}

wb = load_workbook(SOURCE, data_only=False)
replacements = 0
for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                value = cell.value
                updated = value
                for old, new in mapping.items():
                    updated = updated.replace(old, new)
                if updated != value:
                    cell.value = updated
                    replacements += 1

# Make the title explicit without changing plan content.
wb["Tổng quan"]["A1"] = "QUICKASSIST — PHÂN TÍCH & KẾ HOẠCH NHÓM 4 NGƯỜI (BẢN ẨN DANH)"
wb["Backlog phân công"]["A1"] = "BACKLOG PHÂN CÔNG CHI TIẾT — BẢN ẨN DANH"
wb["Timeline 8 tuần"]["A1"] = "TIMELINE 8 TUẦN — BẢN ẨN DANH"
wb["Tải công việc"]["A1"] = "TẢI CÔNG VIỆC THEO THÀNH VIÊN — BẢN ẨN DANH"
wb["Rủi ro & quyết định"]["A1"] = "RỦI RO, KHOẢNG TRỐNG & HƯỚNG XỬ LÝ — BẢN ẨN DANH"
wb["Kế hoạch tuần"]["A1"] = "MỤC TIÊU, MỐC BÀN GIAO & NHỊP LÀM VIỆC — BẢN ẨN DANH"

wb.properties.title = "QuickAssist — Kế hoạch phân công ẩn danh"
wb.properties.subject = "Kế hoạch giữ nguyên phân công, thay tên thành Người A–D"
wb.properties.creator = "Nhóm 1"
wb.save(OUTPUT)

check = load_workbook(OUTPUT, data_only=False)
backlog = check["Backlog phân công"]
owners = {backlog.cell(r, 5).value for r in range(4, backlog.max_row + 1)}
assert owners == {"Người A", "Người B", "Người C", "Người D"}
assert "Mai Đức Vinh" not in "\n".join(
    str(cell.value) for sheet in check.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None
)
assert check["Tải công việc"]["D4"].value.startswith("=SUMIF")
print(f"Created: {OUTPUT}")
print(f"Identity replacements: {replacements}")
print("Owners:", ", ".join(sorted(owners)))
