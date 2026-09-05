from pathlib import Path
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "QuickAssist_KeHoach_PhanCong_AnDanh_GoogleOAuth.xlsx"
OUTPUT = ROOT / "QuickAssist_KeHoach_PhanCong_GoogleOAuth_CoTen.xlsx"

# Keep existing A–D task assignments exactly as they are; only reveal identities.
mapping = {
    "Người A": "Lê Đăng Sơn",
    "Người B": "Trịnh Mạnh Quang",
    "Người C": "Mai Đức Vinh",
    "Người D": "Trần Thùy Dương",
}

wb = load_workbook(SOURCE, data_only=False)
replaced = 0
for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                before = cell.value
                after = before
                for placeholder, name in mapping.items():
                    after = after.replace(placeholder, name)
                if after != before:
                    cell.value = after
                    replaced += 1

wb["Tổng quan"]["A1"] = "QUICKASSIST — PHÂN TÍCH & KẾ HOẠCH NHÓM 4 NGƯỜI"
wb["Backlog phân công"]["A1"] = "BACKLOG PHÂN CÔNG CHI TIẾT"
wb["Timeline 8 tuần"]["A1"] = "TIMELINE 8 TUẦN"
wb["Tải công việc"]["A1"] = "TẢI CÔNG VIỆC THEO THÀNH VIÊN"
wb["Rủi ro & quyết định"]["A1"] = "RỦI RO, KHOẢNG TRỐNG & HƯỚNG XỬ LÝ"
wb["Kế hoạch tuần"]["A1"] = "MỤC TIÊU, MỐC BÀN GIAO & NHỊP LÀM VIỆC"
wb.properties.title = "QuickAssist — Kế hoạch phân công Google OAuth"
wb.properties.subject = "Kế hoạch 8 tuần có phân công theo thành viên"
wb.save(OUTPUT)

check = load_workbook(OUTPUT, data_only=False)
backlog = check["Backlog phân công"]
owners = {backlog.cell(r, 5).value for r in range(4, backlog.max_row + 1)}
assert owners == set(mapping.values())
assert all(name not in "\n".join(str(c.value) for s in check.worksheets for row in s.iter_rows() for c in row if c.value is not None) for name in mapping)
assert check["Tải công việc"]["D4"].value.startswith("=SUMIF")
print(f"Created: {OUTPUT}")
print(f"Replacements: {replaced}")
print("Owners:", ", ".join(sorted(owners)))
