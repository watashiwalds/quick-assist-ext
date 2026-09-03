from copy import copy
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "QuickAssist_KeHoach_PhanCong_4Nguoi.xlsx"
OUTPUT = ROOT / "QuickAssist_KeHoach_TamThoi_ChuaPhanCong.xlsx"

wb = load_workbook(SOURCE)

navy = "17365D"
blue = "2F75B5"
light_blue = "D9EAF7"
gray = "E7E6E6"
light_gray = "F2F2F2"
white = "FFFFFF"

def set_note(sheet, cell_range, text):
    start = sheet[cell_range.split(":")[0]]
    start.value = text
    start.font = Font(bold=True, color=navy, size=10)
    start.fill = PatternFill("solid", fgColor=light_blue)
    start.alignment = Alignment(wrap_text=True, vertical="center")

# Tổng quan: anonymize members and state the plan is intentionally unassigned.
ws = wb["Tổng quan"]
ws["B4"] = "8 tuần — phạm vi và mốc đã xác định; phân vai/ước lượng theo người sẽ chốt sau."
ws["B5"] = "Tài khoản/auth + ghi chú nhanh + thư viện/CRUD/cache + RAG cơ bản + bảo mật/kiểm thử."
ws["B6"] = "Bản tạm: chưa phân công task cho thành viên; Skim là P1, lịch trình cá nhân là P2."
ws["A17"] = "Danh sách tạm"
ws["A18"] = "Thành viên tạm"
ws["B18"] = "Vai trò"
ws["C18"] = "Ghi chú"
for r, label in enumerate(["Người A", "Người B", "Người C", "Người D"], 19):
    ws.cell(r, 1).value = label
    ws.cell(r, 2).value = "Chưa phân vai"
    ws.cell(r, 3).value = "Sẽ xác định sau khi nhóm thống nhất năng lực, thời gian và phần việc."
    for c in range(1, 7):
        ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")

# Backlog: keep scope/timeline, eliminate all individual allocation.
bl = wb["Backlog phân công"]
bl["A1"] = "BACKLOG KẾ HOẠCH TẠM — CHƯA PHÂN CÔNG"
last_row = bl.max_row
for r in range(4, last_row + 1):
    bl.cell(r, 5).value = "Chưa phân công"
    bl.cell(r, 6).value = "—"
    bl.cell(r, 14).value = "Chưa phân công"
    bl.cell(r, 15).value = 0
    for c in (5, 6, 14):
        bl.cell(r, c).fill = PatternFill("solid", fgColor=light_gray)
        bl.cell(r, c).font = Font(italic=True, color="666666")

# Timeline: remove named owner color meaning, retain schedule bars in neutral gray.
tl = wb["Timeline 8 tuần"]
tl["A1"] = "TIMELINE 8 TUẦN — CHƯA GÁN NGƯỜI PHỤ TRÁCH"
for r in range(4, tl.max_row + 1):
    tl.cell(r, 3).value = "Chưa phân công"
    tl.cell(r, 3).fill = PatternFill("solid", fgColor=light_gray)
    tl.cell(r, 3).font = Font(italic=True, color="666666")
    for c in range(6, 14):
        cell = tl.cell(r, c)
        if cell.value == "■":
            cell.fill = PatternFill("solid", fgColor="BFBFBF")
            cell.font = Font(color="595959", bold=True)

# Workload: reset to placeholders, no owner formulas or chart suggesting allocation.
wl = wb["Tải công việc"]
wl["A1"] = "TẢI CÔNG VIỆC — BẢN TẠM CHƯA PHÂN CÔNG"
wl["A3"] = "Thành viên tạm"
wl["B3"] = "Vai trò"
wl["C3"] = "Trạng thái"
wl["D3"] = "Task đã giao"
for c in range(5, 13):
    wl.cell(3, c).value = "—"
for r, label in enumerate(["Người A", "Người B", "Người C", "Người D"], 4):
    wl.cell(r, 1).value = label
    wl.cell(r, 2).value = "Chưa phân vai"
    wl.cell(r, 3).value = "Chưa phân công"
    wl.cell(r, 4).value = 0
    for c in range(5, 13):
        wl.cell(r, c).value = "—"
    for c in range(1, 13):
        wl.cell(r, c).fill = PatternFill("solid", fgColor=light_gray)
        wl.cell(r, c).alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if c >= 3 else "left")
wl["A10"] = "Trạng thái"
wl["B10"] = "Chưa gán task cho Người A–D. Khi nhóm chốt vai trò, điền Owner trong sheet “Backlog phân công”; tải công việc sẽ được thiết lập lại theo phân công chính thức."
if len(wl._charts):
    wl._charts = []

# Risks and weekly reviews must not imply individual allocation.
rk = wb["Rủi ro & quyết định"]
rk["A1"] = "RỦI RO, KHOẢNG TRỐNG & HƯỚNG XỬ LÝ — OWNER CHƯA PHÂN CÔNG"
for r in range(4, rk.max_row + 1):
    rk.cell(r, 6).value = "Chưa phân công"
    rk.cell(r, 6).fill = PatternFill("solid", fgColor=light_gray)
    rk.cell(r, 6).font = Font(italic=True, color="666666")

wk = wb["Kế hoạch tuần"]
wk["A1"] = "MỤC TIÊU, MỐC BÀN GIAO & NHỊP LÀM VIỆC — CHƯA PHÂN VAI"
for r in range(4, 12):
    wk.cell(r, 5).value = "Cả nhóm (chưa phân vai)"
wk["B14"] = "Đầu tuần: chốt mục tiêu & dependency (30’). Giữa tuần: sync blocker (15’). Cuối tuần: demo + review + cập nhật backlog (45’). Việc gán owner và reviewer sẽ thực hiện sau khi nhóm chốt phân vai."

# Ensure workbook metadata identifies template status.
wb.properties.title = "QuickAssist — Kế hoạch tạm chưa phân công"
wb.properties.subject = "Kế hoạch 8 tuần và backlog chưa gán thành viên"
wb.properties.creator = "Nhóm 1"
wb.properties.keywords = "QuickAssist, kế hoạch, template, chưa phân công"

wb.save(OUTPUT)

# Validate the transformation.
check = load_workbook(OUTPUT, data_only=False)
assert check["Backlog phân công"]["E4"].value == "Chưa phân công"
assert check["Backlog phân công"]["N4"].value == "Chưa phân công"
assert check["Tải công việc"]["A4"].value == "Người A"
assert check["Tải công việc"]["D4"].value == 0
assert all(check["Rủi ro & quyết định"].cell(r, 6).value == "Chưa phân công" for r in range(4, check["Rủi ro & quyết định"].max_row + 1))
print(f"Created: {OUTPUT}")
print(f"Backlog rows marked unassigned: {last_row - 3}")
