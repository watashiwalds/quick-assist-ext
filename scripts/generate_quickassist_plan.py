from pathlib import Path
from collections import defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, ColorScaleRule
from openpyxl.chart import BarChart, Reference


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "QuickAssist_KeHoach_PhanCong_4Nguoi.xlsx"

members = [
    ("Mai Đức Vinh", "Tech Lead / Browser Extension", "Kiến trúc tổng thể, Manifest V3, TypeScript, IndexedDB, tích hợp UI–API"),
    ("Trịnh Mạnh Quang", "Backend Lead", "API Python, xác thực, PostgreSQL/pgvector, quota, đồng bộ dữ liệu"),
    ("Lê Đăng Sơn", "AI/RAG Lead", "Chunking, embedding, hybrid search, reranking, đánh giá chất lượng AI"),
    ("Trần Thuỳ Dương", "UI/UX & QA Lead", "Luồng giao diện, trạng thái lỗi, kiểm thử, tài liệu và demo"),
]

tasks = [
    ("P01", "Khởi động", "Kiến trúc", "Chốt kiến trúc, module, luồng dữ liệu và ADR", members[0][0], members[1][0], 1, 1, 6, "P0", "", "Sơ đồ kiến trúc + ADR", "Nhóm duyệt; thể hiện extension/server/AI/DB và ranh giới dữ liệu"),
    ("P02", "Khởi động", "API & dữ liệu", "Chuẩn hoá endpoint, Google OAuth callback/session và data contract; xử lý xung đột /notes/quick, /notes, /notes/ragged", members[1][0], members[0][0], 1, 1, 8, "P0", "P01", "OpenAPI draft + schema", "Không còn endpoint mâu thuẫn; có OAuth flow, request/response/error mẫu"),
    ("P03", "Khởi động", "AI", "Spike lựa chọn AI Provider: Dify hay service riêng; đo khả năng chunk/embed/rerank", members[2][0], members[1][0], 1, 1, 8, "P0", "P01", "Báo cáo spike + quyết định", "Có tiêu chí latency, chi phí/quota, khả năng chạy local"),
    ("P04", "Khởi động", "UX", "Vẽ user flow và wireframe cho 5 chức năng trong kịch bản", members[3][0], members[0][0], 1, 1, 8, "P0", "P01", "Wireframe + user flow", "Bao phủ happy path, loading, empty, error, undo"),
    ("P05", "Nền tảng", "Extension", "Khởi tạo TypeScript + Manifest V3, build/lint/test và cấu trúc module", members[0][0], members[3][0], 1, 2, 8, "P0", "P01", "Extension skeleton", "Cài được ở developer mode; build và lint pass"),
    ("P06", "Nền tảng", "Backend", "Khởi tạo Python API, cấu hình môi trường, health check và test framework", members[1][0], members[2][0], 1, 2, 8, "P0", "P01", "Backend skeleton", "Health check chạy; test smoke pass; cấu hình tách dev/test"),
    ("P07", "Nền tảng", "AI", "Khởi tạo AI service/adapter cho local provider và mock provider", members[2][0], members[1][0], 1, 2, 8, "P0", "P03,P06", "AI adapter skeleton", "Đổi provider qua config; có mock deterministic cho test"),
    ("P08", "Nền tảng", "QA", "Lập test strategy, traceability matrix từ kịch bản và Definition of Done", members[3][0], members[0][0], 1, 2, 6, "P0", "P04", "Test plan + RTM", "Mỗi nhánh chính có test case và người chịu trách nhiệm"),
    ("A01", "Nền tảng", "Tài khoản", "Thiết kế schema user/note/chunk/category/quota và migration", members[1][0], members[2][0], 2, 2, 8, "P0", "P02,P06", "Migration + ERD", "Có UUID, ownership, timestamp, index và ràng buộc dữ liệu"),
    ("A02", "Nền tảng", "Tài khoản", "Xây Google OAuth 2.0 Authorization Code + PKCE: callback, verify ID token và phát hành app session", members[1][0], members[0][0], 2, 3, 10, "P0", "A01", "Google OAuth API", "Chỉ dùng scope openid/email/profile; validate state, nonce, issuer, audience và expiry"),
    ("A03", "Nền tảng", "Tài khoản", "Tích hợp Google OAuth trong extension qua browser identity flow/PKCE và quản lý app session", members[0][0], members[1][0], 2, 3, 8, "P0", "A02,P05", "Google OAuth UI + client", "Login/logout hoạt động; không log token; xử lý session hết hạn; không có form mật khẩu"),
    ("A04", "Nền tảng", "Quota", "Xây dựng quota middleware và nhật ký tiêu hao AI", members[1][0], members[2][0], 3, 3, 8, "P0", "A01,A02", "Quota service", "Chặn đúng khi hết quota; cập nhật nguyên tử; có audit"),
    ("A05", "Nền tảng", "UI", "Xây design system và app shell: loading/empty/error/toast/modal", members[3][0], members[0][0], 2, 3, 8, "P0", "P04,P05", "UI component set", "Các màn hình dùng chung component; keyboard usable"),
    ("N01", "MVP Ghi chú", "Extension", "Bắt selection, context menu và phím Ctrl-Alt-N", members[0][0], members[3][0], 3, 3, 8, "P0", "P05", "Selection capture", "Lấy text + URL + title; hoạt động trên 3 trang mẫu"),
    ("N02", "MVP Ghi chú", "UI", "Popup snippet, đếm ngược 5 giây và hoàn tác Ctrl-Alt-Z", members[3][0], members[0][0], 3, 4, 8, "P0", "A05,N01", "Undo popup", "Undo hủy note mới nhất; trạng thái hiển thị rõ"),
    ("N03", "MVP Ghi chú", "Backend", "Cài API tạo note nhanh, xác thực ownership và lưu PostgreSQL", members[1][0], members[0][0], 3, 4, 10, "P0", "A01,A02,P02", "POST note API", "Validation/idempotency/error contract và integration test pass"),
    ("N04", "MVP Ghi chú", "AI", "Cài semantic chunking + embedding, định danh chunk UUID", members[2][0], members[1][0], 3, 4, 10, "P0", "P07", "Chunk/embed pipeline", "Giữ liên kết note–chunk; xử lý text dài và lỗi provider"),
    ("N05", "MVP Ghi chú", "AI/Backend", "Lưu vector pgvector và cập nhật quota sau xử lý", members[2][0], members[1][0], 4, 4, 8, "P0", "N03,N04,A04", "Vector persistence", "Transaction nhất quán; không trừ quota khi silent skip"),
    ("N06", "MVP Ghi chú", "Extension", "Local state cho hàng chờ, retry và trạng thái note", members[0][0], members[1][0], 4, 4, 8, "P0", "N02,N03", "Queue/state store", "Không tạo note trùng khi retry; hiển thị pending/success/error"),
    ("N07", "MVP Ghi chú", "QA", "Kiểm thử E2E ghi chú nhanh, undo, hết quota và mất mạng", members[3][0], members[0][0], 4, 4, 6, "P0", "N02,N03,N04,N05,N06", "E2E suite + biên bản", "Các ca P0 pass; lỗi có bằng chứng và owner"),
    ("C01", "Thư viện note", "Backend", "API GET/PUT/DELETE note, filter metadata/category và bulk delete", members[1][0], members[0][0], 4, 5, 10, "P0", "N03", "CRUD API", "Ownership bắt buộc; filter/pagination; bulk delete an toàn"),
    ("C02", "Thư viện note", "Extension", "IndexedDB cache/offline storage và chiến lược sync", members[0][0], members[1][0], 4, 5, 10, "P0", "N06,C01", "Cache/sync module", "Cache-first; sync theo updatedAt; có quy tắc xử lý xung đột"),
    ("C03", "Thư viện note", "UI", "Màn hình danh sách/chi tiết/sửa/xoá note và phân mục", members[3][0], members[0][0], 4, 5, 10, "P0", "A05,C01", "Note library UI", "Browse/edit/delete/filter chạy với cache và server"),
    ("C04", "Thư viện note", "Extension", "Tích hợp metadata, category filter và đồng bộ trạng thái UI", members[0][0], members[3][0], 5, 5, 6, "P0", "C02,C03", "Filter integration", "Kết quả đúng giữa local/server; refresh không mất trạng thái"),
    ("C05", "Thư viện note", "QA", "Test CRUD, cache-first, offline, sync conflict và xoá hàng loạt", members[3][0], members[1][0], 5, 5, 6, "P0", "C01,C02,C03,C04", "CRUD/offline test suite", "Ca dữ liệu trống/lỗi mạng/xung đột đều được xác minh"),
    ("R01", "RAG", "AI", "Xây hybrid retrieval: BM25 + cosine similarity + search pool theo user", members[2][0], members[1][0], 4, 5, 12, "P0", "N04,N05", "Hybrid retrieval", "Không rò dữ liệu chéo user; có top-k và score"),
    ("R02", "RAG", "AI", "Tích hợp rerank top-k và fallback khi provider lỗi", members[2][0], members[1][0], 5, 5, 8, "P0", "R01,P07", "Rerank adapter", "Trả top-10 ổn định; timeout/fallback được test"),
    ("R03", "RAG", "Backend", "Cài POST /search/ranking: auth, quota, metadata filter và response", members[1][0], members[2][0], 5, 6, 10, "P0", "R01,R02,A04", "Search API", "Contract đúng; quota đúng; filter category pass"),
    ("R04", "RAG", "Extension", "Màn hình tìm kiếm/chat, gọi API và render chunk/rank", members[0][0], members[3][0], 5, 6, 8, "P0", "A05,R03", "RAG search UI", "Loading/error/empty/result rõ; mở được note nguồn"),
    ("R05", "RAG", "QA/UX", "Đánh giá trải nghiệm kết quả RAG và test phân quyền/quota", members[3][0], members[2][0], 6, 6, 6, "P0", "R03,R04", "RAG acceptance report", "Không lộ dữ liệu; kết quả đáp ứng bộ query mẫu"),
    ("S01", "Skim trang", "Extension", "Readability extraction, session UUID theo tab+user", members[0][0], members[2][0], 6, 6, 8, "P1", "P05", "Page extraction", "Loại bỏ menu/quảng cáo; session tách theo tab/user"),
    ("S02", "Skim trang", "AI/Backend", "Cài /skim/prepare, /skim/search và bộ nhớ session có TTL", members[2][0], members[1][0], 6, 7, 10, "P1", "S01,R01,R02", "Skim APIs", "Có TTL/limit; xử lý session missing/expired"),
    ("S03", "Skim trang", "AI", "Pipeline chunk/embed/search/rerank tạm thời cho nội dung trang", members[2][0], members[1][0], 6, 7, 10, "P1", "S01,R01,R02", "Skim AI pipeline", "Không ghi lâu dài trước khi user chọn; latency được đo"),
    ("S04", "Skim trang", "UI", "UI prompt, hint chuẩn bị/tìm kiếm, preview và chọn đoạn để lưu", members[3][0], members[0][0], 6, 7, 8, "P1", "A05,S01,S02", "Skim UI", "Đúng trạng thái Ready/chưa Ready; chọn nhiều đoạn được"),
    ("S05", "Skim trang", "Backend", "Cài lưu các chunk đã chọn qua contract note thống nhất", members[1][0], members[0][0], 7, 7, 6, "P1", "S02,S03,P02", "Save selected chunks", "Lưu note+vector đúng ownership; không tạo trùng"),
    ("S06", "Skim trang", "QA", "E2E skim: trang dài, chưa ready, hết quota, session hết hạn", members[3][0], members[2][0], 7, 7, 6, "P1", "S02,S03,S04,S05", "Skim E2E report", "Happy path và nhánh lỗi trọng yếu pass"),
    ("G01", "Khoảng trống", "Đặc tả", "Đặc tả riêng chức năng Tóm tắt nội dung đang thiếu trong file kịch bản", members[3][0], members[2][0], 2, 2, 5, "P1", "P04", "Scenario + acceptance", "Chốt đầu vào/đầu ra/lưu trữ/quota/privacy; quyết định nhận vào P1 hoặc hoãn"),
    ("H01", "Hoàn thiện", "Bảo mật", "Rà soát permission MV3, CSP, secret/token và dữ liệu nhạy cảm", members[0][0], members[1][0], 7, 8, 6, "P0", "A03,C02,S01", "Security checklist", "Least privilege; không secret trong bundle/log"),
    ("H02", "Hoàn thiện", "Bảo mật", "Validation, Google OAuth token/session validation, rate limit, CORS/HTTPS config và kiểm tra ownership API", members[1][0], members[0][0], 7, 8, 8, "P0", "C01,R03,S02,A02", "Backend hardening", "OAuth security tests pass; lỗi không lộ chi tiết nội bộ"),
    ("H03", "Hoàn thiện", "AI", "Tạo bộ dữ liệu đánh giá và benchmark relevance/latency", members[2][0], members[3][0], 7, 8, 8, "P0", "R01,R02,S03", "Evaluation report", "Có Recall@k/nDCG hoặc tiêu chí tương đương và ngưỡng chấp nhận"),
    ("H04", "Hoàn thiện", "AI/Deploy", "Cấu hình self-host/Dify, health/fallback và tài liệu vận hành", members[2][0], members[1][0], 7, 8, 8, "P1", "P03,P07", "AI deployment guide", "Một cấu hình demo chạy được; provider lỗi có cảnh báo"),
    ("H05", "Hoàn thiện", "QA", "Regression E2E, usability, accessibility và kiểm thử demo", members[3][0], members[0][0], 8, 8, 10, "P0", "N07,C05,R05,S06,H01,H02", "Release test report", "Không còn lỗi P0/P1 mở; checklist demo pass"),
    ("H06", "Hoàn thiện", "Tích hợp", "Tối ưu hiệu năng, retry/telemetry và sửa lỗi tích hợp extension", members[0][0], members[3][0], 8, 8, 8, "P0", "R04,S04,H01", "Release candidate extension", "Build sạch; luồng P0 ổn định; lỗi có log không chứa dữ liệu nhạy cảm"),
    ("H07", "Hoàn thiện", "Tích hợp", "Đóng gói backend, migration, cấu hình demo và sửa lỗi tích hợp", members[1][0], members[2][0], 8, 8, 8, "P0", "H02,H04", "Release candidate server", "Deploy mới từ hướng dẫn; migration và smoke test pass"),
    ("H08", "Hoàn thiện", "Tài liệu", "Hoàn thiện báo cáo, hướng dẫn người dùng và kịch bản thuyết trình", members[3][0], members[0][0], 8, 8, 8, "P0", "H03,H05,H06,H07", "Report + user guide + demo script", "Tài liệu khớp sản phẩm; demo 10–15 phút chạy được"),
]

risks = [
    ("R-01", "Tóm tắt nội dung có trong Word nhưng chưa có kịch bản chi tiết", "Cao", "Cao", "Đặc tả ở G01; chỉ đưa vào MVP sau khi chốt acceptance/quota", members[3][0]),
    ("R-02", "Tên endpoint ghi chú không thống nhất giữa các bước", "Cao", "Trung bình", "Chuẩn hoá OpenAPI ở P02 trước khi code client/server", members[1][0]),
    ("R-03", "Google OAuth cần xác thực callback, state/nonce và app session rõ ràng", "Cao", "Cao", "Authorization Code + PKCE; validate issuer/audience/expiry; ownership check mọi truy vấn", members[1][0]),
    ("R-04", "Quota và silent skip AI có thể gây trạng thái note khó hiểu", "Trung bình", "Cao", "Tách trạng thái saved/processed; log lý do và hiển thị phù hợp", members[0][0]),
    ("R-05", "Dữ liệu web cá nhân được gửi đến AI provider", "Cao", "Trung bình", "Consent, tối thiểu hoá dữ liệu, retention policy và tùy chọn local", members[0][0]),
    ("R-06", "Hybrid search/rerank có độ trễ cao", "Trung bình", "Cao", "Benchmark sớm, timeout, cache và fallback không rerank", members[2][0]),
    ("R-07", "IndexedDB và server có thể xung đột dữ liệu", "Trung bình", "Cao", "Chốt version/updatedAt và quy tắc conflict trong C02", members[0][0]),
    ("R-08", "Readability không ổn định trên SPA/trang đặc biệt", "Trung bình", "Trung bình", "Bộ trang thử chuẩn, fallback selection/body text và thông báo giới hạn", members[0][0]),
    ("R-09", "Self-host LLM vượt tài nguyên máy demo", "Cao", "Trung bình", "Provider adapter, model nhẹ, mock/demo fallback", members[2][0]),
    ("R-10", "Phạm vi rộng so với nhóm 4 người", "Cao", "Cao", "Khoá P0; Skim và tóm tắt chỉ nhận khi P0 ổn định", members[0][0]),
]

features = [
    ("Tài khoản cá nhân", "Word", "Một phần", "P0", "Google OAuth 2.0 + app session, quota và quyền riêng tư"),
    ("Ghi chú nhanh", "Word + Excel", "Chi tiết", "P0", "Luồng selection → undo 5s → lưu → AI xử lý → local state"),
    ("Tìm văn bản nổi bật trên trang", "Excel", "Chi tiết", "P1", "Readability + session tạm + hybrid search + chọn đoạn để lưu"),
    ("Duyệt note/cache", "Excel", "Chi tiết", "P0", "Cache-first IndexedDB, sync server và hiển thị thư viện"),
    ("Tra cứu note RAG-like", "Word + Excel", "Chi tiết", "P0", "Filter theo user/category; BM25 + vector + rerank"),
    ("Note CRUD", "Excel", "Chi tiết", "P0", "Sửa/xoá/bulk delete và cập nhật local tương ứng"),
    ("Tóm tắt nội dung", "Word", "Thiếu kịch bản", "P1", "Phải bổ sung scenario, API, UI, quota, lưu trữ và acceptance"),
    ("Tự triển khai/local AI", "Word", "Định hướng", "P1", "Cần adapter và hướng dẫn cấu hình, không khóa vào một provider"),
]

wb = Workbook()
ws = wb.active
ws.title = "Tổng quan"

navy = "17365D"
blue = "2F75B5"
light_blue = "D9EAF7"
teal = "008C95"
light_teal = "DDEBF7"
orange = "F4B183"
light_orange = "FCE4D6"
green = "70AD47"
light_green = "E2F0D9"
red = "C00000"
light_red = "F4CCCC"
gray = "E7E6E6"
dark_gray = "595959"
white = "FFFFFF"
thin = Side(style="thin", color="B7B7B7")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def title(sheet, text, end_col):
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    c = sheet.cell(1, 1, text)
    c.font = Font(bold=True, color=white, size=18)
    c.fill = PatternFill("solid", fgColor=navy)
    c.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30

def header_row(sheet, row, cols):
    for col in range(1, cols + 1):
        c = sheet.cell(row, col)
        c.font = Font(bold=True, color=white)
        c.fill = PatternFill("solid", fgColor=blue)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border

def style_body(sheet, min_row, max_row, min_col, max_col):
    for row in sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for c in row:
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=True)

title(ws, "QUICKASSIST — PHÂN TÍCH & KẾ HOẠCH NHÓM 4 NGƯỜI", 6)
ws["A3"] = "Mục tiêu"
ws["B3"] = "Xây extension ghi chú nhanh, quản lý note và tra cứu ngữ nghĩa; kiến trúc mở để dùng dịch vụ hoặc tự triển khai."
ws["A4"] = "Mốc kế hoạch"
ws["B4"] = "8 tuần — khóa phạm vi P0 trước; P1 chỉ làm sau khi P0 đạt Definition of Done."
ws["A5"] = "MVP đề xuất"
ws["B5"] = "Google OAuth + ghi chú nhanh + thư viện/CRUD/cache + RAG cơ bản + bảo mật/kiểm thử."
ws["A6"] = "Ngoài MVP"
ws["B6"] = "Skim trang là P1; tóm tắt phải bổ sung kịch bản trước khi triển khai."
ws["A7"] = "Nguồn phân tích"
ws["B7"] = "Nhóm 1 - Báo cáo ý tưởng đề tài.docx; Kịch bản chức năng QuickAssist_1.xlsx"
for r in range(3, 8):
    ws.cell(r, 1).font = Font(bold=True, color=navy)
    ws.cell(r, 1).fill = PatternFill("solid", fgColor=light_blue)
    ws.cell(r, 1).border = border
    ws.cell(r, 2).border = border
    ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("B3:F3"); ws.merge_cells("B4:F4"); ws.merge_cells("B5:F5"); ws.merge_cells("B6:F6"); ws.merge_cells("B7:F7")

ws["A9"] = "Nhận định chính"
ws["A9"].font = Font(bold=True, color=white, size=13)
ws["A9"].fill = PatternFill("solid", fgColor=teal)
ws.merge_cells("A9:F9")
insights = [
    "Kiến trúc gồm Extension/UI ↔ Server ↔ AI Provider, với PostgreSQL + pgvector và IndexedDB làm cache/offline.",
    "Kịch bản đã mô tả tốt 5 luồng: ghi chú nhanh, skim trang, duyệt note, RAG và CRUD; nhưng account/auth mới ở mức ý tưởng.",
    "Mấu chốt dữ liệu là ownership theo user, metadata/category, UUID note/chunk, quota và đồng bộ local–server.",
    "Hai contract phải chốt trước code: endpoint ghi chú và quy tắc response/error/idempotency.",
    "Tóm tắt nội dung chưa đủ đặc tả; không nên cam kết trong MVP trước khi hoàn thành G01.",
    "RAG nên có benchmark chất lượng/độ trễ; luôn có timeout/fallback để demo không phụ thuộc tuyệt đối vào provider.",
]
for i, item in enumerate(insights, 10):
    ws.cell(i, 1, f"{i-9}.")
    ws.cell(i, 2, item)
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
    for c in range(1, 7):
        ws.cell(i, c).border = border
        ws.cell(i, c).alignment = Alignment(wrap_text=True, vertical="top")

ws["A18"] = "Thành viên"
ws["B18"] = "Vai trò chính"
ws["C18"] = "Trách nhiệm"
ws.merge_cells("C18:F18")
header_row(ws, 18, 6)
for r, member in enumerate(members, 19):
    ws.cell(r, 1, member[0]); ws.cell(r, 2, member[1]); ws.cell(r, 3, member[2])
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    for c in range(1, 7):
        ws.cell(r, c).border = border; ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")

ws["A25"] = "Definition of Done chung"
ws["A25"].font = Font(bold=True, color=white, size=13); ws["A25"].fill = PatternFill("solid", fgColor=teal); ws.merge_cells("A25:F25")
dod = [
    "Có acceptance criteria và test tương ứng; code được thành viên phối hợp review.",
    "Build/lint/unit/integration liên quan đều pass; không còn lỗi P0/P1 mở.",
    "Không log token/nội dung nhạy cảm; kiểm tra ownership và validation ở server.",
    "API/schema/tài liệu được cập nhật cùng thay đổi; demo được trên dữ liệu mẫu.",
]
for i, item in enumerate(dod, 26):
    ws.cell(i, 1, "✓"); ws.cell(i, 2, item); ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
    for c in range(1, 7): ws.cell(i, c).border = border; ws.cell(i, c).alignment = Alignment(wrap_text=True, vertical="top")

ws.column_dimensions["A"].width = 21; ws.column_dimensions["B"].width = 28
for col in "CDEF": ws.column_dimensions[col].width = 18
ws.freeze_panes = "A3"

# Feature matrix
fm = wb.create_sheet("Ma trận chức năng")
title(fm, "MA TRẬN PHẠM VI & MỨC ĐỘ ĐẶC TẢ", 5)
headers = ["Chức năng", "Nguồn", "Mức đặc tả", "Ưu tiên", "Nhận xét / quyết định"]
for c, h in enumerate(headers, 1): fm.cell(3, c, h)
header_row(fm, 3, len(headers))
for r, row in enumerate(features, 4):
    for c, v in enumerate(row, 1): fm.cell(r, c, v)
style_body(fm, 4, 3 + len(features), 1, 5)
fm.auto_filter.ref = f"A3:E{3+len(features)}"; fm.freeze_panes = "A4"
for col, width in zip("ABCDE", [30, 18, 18, 12, 65]): fm.column_dimensions[col].width = width

# Backlog
bl = wb.create_sheet("Backlog phân công")
title(bl, "BACKLOG PHÂN CÔNG CHI TIẾT", 15)
headers = ["ID", "Giai đoạn", "Hạng mục", "Task", "Người phụ trách", "Phối hợp", "Tuần BĐ", "Tuần KT", "Giờ", "Ưu tiên", "Phụ thuộc", "Deliverable", "Tiêu chí nghiệm thu", "Trạng thái", "% hoàn thành"]
for c, h in enumerate(headers, 1): bl.cell(3, c, h)
header_row(bl, 3, len(headers))
for r, t in enumerate(tasks, 4):
    values = list(t) + ["Chưa bắt đầu", 0]
    for c, v in enumerate(values, 1): bl.cell(r, c, v)
    bl.cell(r, 15).number_format = "0%"
style_body(bl, 4, 3 + len(tasks), 1, 15)
bl.auto_filter.ref = f"A3:O{3+len(tasks)}"; bl.freeze_panes = "A4"
widths = [9, 15, 16, 48, 22, 22, 9, 9, 8, 9, 18, 28, 48, 17, 14]
for i, width in enumerate(widths, 1): bl.column_dimensions[chr(64+i)].width = width
dv_status = DataValidation(type="list", formula1='"Chưa bắt đầu,Đang làm,Chờ review,Bị chặn,Hoàn thành"', allow_blank=False)
dv_progress = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=False)
bl.add_data_validation(dv_status); bl.add_data_validation(dv_progress)
dv_status.add(f"N4:N{3+len(tasks)}"); dv_progress.add(f"O4:O{3+len(tasks)}")
bl.conditional_formatting.add(f"O4:O{3+len(tasks)}", ColorScaleRule(start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=.5, mid_color="FFEB84", end_type="num", end_value=1, end_color="63BE7B"))
bl.conditional_formatting.add(f"N4:N{3+len(tasks)}", FormulaRule(formula=["$N4=\"Bị chặn\""], fill=PatternFill("solid", fgColor=light_red)))

# Timeline
tl = wb.create_sheet("Timeline 8 tuần")
title(tl, "TIMELINE 8 TUẦN", 13)
headers = ["ID", "Task", "Owner", "Ưu tiên", "Giờ"] + [f"Tuần {i}" for i in range(1, 9)]
for c, h in enumerate(headers, 1): tl.cell(3, c, h)
header_row(tl, 3, len(headers))
owner_colors = {members[0][0]: "D9EAF7", members[1][0]: "E2F0D9", members[2][0]: "FFF2CC", members[3][0]: "FCE4D6"}
for r, t in enumerate(tasks, 4):
    tl.cell(r, 1, t[0]); tl.cell(r, 2, t[3]); tl.cell(r, 3, t[4]); tl.cell(r, 4, t[9]); tl.cell(r, 5, t[8])
    for week in range(1, 9):
        c = tl.cell(r, 5 + week)
        if t[6] <= week <= t[7]:
            c.value = "■"
            c.fill = PatternFill("solid", fgColor=owner_colors[t[4]])
            c.font = Font(color=navy, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
    for c in range(1, 14): tl.cell(r, c).border = border; tl.cell(r, c).alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if c != 2 else "left")
tl.auto_filter.ref = f"A3:M{3+len(tasks)}"; tl.freeze_panes = "F4"
for col, width in zip(["A","B","C","D","E"], [9, 52, 22, 10, 8]): tl.column_dimensions[col].width = width
for col in "FGHIJKLM": tl.column_dimensions[col].width = 11

# Workload summary with formulas
wl = wb.create_sheet("Tải công việc")
title(wl, "TẢI CÔNG VIỆC THEO THÀNH VIÊN", 12)
headers = ["Thành viên", "Vai trò", "Số task", "Tổng giờ"] + [f"T{i}" for i in range(1, 9)]
for c, h in enumerate(headers, 1): wl.cell(3, c, h)
header_row(wl, 3, len(headers))
last_task_row = 3 + len(tasks)
for r, member in enumerate(members, 4):
    wl.cell(r, 1, member[0]); wl.cell(r, 2, member[1])
    wl.cell(r, 3, f'=COUNTIF(\'Backlog phân công\'!$E$4:$E${last_task_row},A{r})')
    wl.cell(r, 4, f'=SUMIF(\'Backlog phân công\'!$E$4:$E${last_task_row},A{r},\'Backlog phân công\'!$I$4:$I${last_task_row})')
    for week in range(1, 9):
        col = 4 + week
        formula = f'=SUMPRODUCT((\'Backlog phân công\'!$E$4:$E${last_task_row}=$A{r})*(\'Backlog phân công\'!$G$4:$G${last_task_row}<={week})*(\'Backlog phân công\'!$H$4:$H${last_task_row}>={week})*(\'Backlog phân công\'!$I$4:$I${last_task_row}/(\'Backlog phân công\'!$H$4:$H${last_task_row}-\'Backlog phân công\'!$G$4:$G${last_task_row}+1)))'
        wl.cell(r, col, formula)
        wl.cell(r, col).number_format = "0.0"
style_body(wl, 4, 7, 1, 12)
for col, width in zip(["A","B","C","D"], [24, 29, 12, 12]): wl.column_dimensions[col].width = width
for col in "EFGHIJKL": wl.column_dimensions[col].width = 11
wl["A10"] = "Cách đọc"
wl["B10"] = "T1–T8 phân bổ giờ ước tính đều theo số tuần của task. Dùng để nhận biết chênh lệch tải; không phải timesheet thực tế."
wl.merge_cells("B10:L10")
for c in range(1, 13): wl.cell(10, c).border = border; wl.cell(10, c).alignment = Alignment(wrap_text=True)
wl["A10"].font = Font(bold=True, color=navy); wl["A10"].fill = PatternFill("solid", fgColor=light_blue)
chart = BarChart(); chart.type = "col"; chart.style = 10; chart.title = "Tổng giờ dự kiến"; chart.y_axis.title = "Giờ"; chart.x_axis.title = "Thành viên"
chart.add_data(Reference(wl, min_col=4, min_row=3, max_row=7), titles_from_data=True)
chart.set_categories(Reference(wl, min_col=1, min_row=4, max_row=7)); chart.height = 7; chart.width = 14
wl.add_chart(chart, "A13")

# Risks
rk = wb.create_sheet("Rủi ro & quyết định")
title(rk, "RỦI RO, KHOẢNG TRỐNG & HƯỚNG XỬ LÝ", 6)
headers = ["ID", "Rủi ro / khoảng trống", "Ảnh hưởng", "Khả năng", "Giảm thiểu / quyết định", "Owner"]
for c, h in enumerate(headers, 1): rk.cell(3, c, h)
header_row(rk, 3, len(headers))
for r, row in enumerate(risks, 4):
    for c, v in enumerate(row, 1): rk.cell(r, c, v)
style_body(rk, 4, 3 + len(risks), 1, 6)
rk.auto_filter.ref = f"A3:F{3+len(risks)}"; rk.freeze_panes = "A4"
for col, width in zip("ABCDEF", [10, 50, 14, 14, 65, 23]): rk.column_dimensions[col].width = width
for r in range(4, 4 + len(risks)):
    if rk.cell(r, 3).value == "Cao": rk.cell(r, 3).fill = PatternFill("solid", fgColor=light_red)
    if rk.cell(r, 4).value == "Cao": rk.cell(r, 4).fill = PatternFill("solid", fgColor=light_orange)

# Weekly ceremonies and milestones
wk = wb.create_sheet("Kế hoạch tuần")
title(wk, "MỤC TIÊU, MỐC BÀN GIAO & NHỊP LÀM VIỆC", 6)
headers = ["Tuần", "Mục tiêu", "Mốc bàn giao", "Điều kiện qua cổng", "Review chính", "Ghi chú"]
for c, h in enumerate(headers, 1): wk.cell(3, c, h)
header_row(wk, 3, len(headers))
weekly = [
    (1, "Chốt kiến trúc, API, AI và UX", "ADR, OpenAPI draft, wireframe, repo skeleton", "Nhóm thống nhất scope P0/P1", "Vinh + cả nhóm", "Không code lệch contract trước P02"),
    (2, "Hoàn thiện nền tảng", "Schema, Google OAuth, test plan, UI shell, đặc tả summary", "Skeleton chạy và migration pass", "Quang/Dương", "Chốt G01"),
    (3, "Auth/quota và ghi chú đầu-cuối", "Auth + selection + popup + API note", "Lưu được note không AI", "Vinh/Quang", "Ưu tiên vertical slice"),
    (4, "MVP ghi chú + vector", "Chunk/embed/vector + queue + E2E", "Luồng N01–N07 pass", "Sơn/Dương", "Mốc MVP-1"),
    (5, "Thư viện note + retrieval", "CRUD/cache/UI + hybrid/rerank", "Offline cơ bản và search backend pass", "Cả nhóm", "Không mở P1 nếu P0 đỏ"),
    (6, "RAG hoàn chỉnh, bắt đầu skim", "Search UI/API + Readability + skim skeleton", "RAG P0 nghiệm thu", "Sơn/Vinh", "Mốc MVP-2"),
    (7, "Skim và hardening", "Skim E2E + security + deploy/eval draft", "P1 pass hoặc cắt khỏi demo", "Quang/Dương", "Khoá tính năng cuối tuần"),
    (8, "Release và thuyết trình", "Regression, release candidate, report, demo", "0 lỗi P0/P1; demo rehearsal pass", "Cả nhóm", "Không thêm tính năng mới"),
]
for r, row in enumerate(weekly, 4):
    for c, v in enumerate(row, 1): wk.cell(r, c, v)
style_body(wk, 4, 11, 1, 6)
for col, width in zip("ABCDEF", [10, 36, 46, 42, 22, 34]): wk.column_dimensions[col].width = width
wk.freeze_panes = "A4"
wk["A14"] = "Nhịp cố định"
wk["B14"] = "Đầu tuần: chốt mục tiêu & dependency (30’). Giữa tuần: sync blocker (15’). Cuối tuần: demo + review + cập nhật backlog (45’). Pull request cần 1 reviewer; thay đổi contract cần owner Backend + client liên quan duyệt."
wk.merge_cells("B14:F14")
for c in range(1, 7): wk.cell(14, c).border = border; wk.cell(14, c).alignment = Alignment(wrap_text=True, vertical="top")
wk["A14"].font = Font(bold=True, color=navy); wk["A14"].fill = PatternFill("solid", fgColor=light_blue)

# Consistent print/view settings
for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddFooter.center.text = "QuickAssist — Kế hoạch nhóm 4 người"

wb.save(OUTPUT)

# Verification: reopen and validate formulas/sheets/key rows.
check = load_workbook(OUTPUT, data_only=False)
assert check.sheetnames == ["Tổng quan", "Ma trận chức năng", "Backlog phân công", "Timeline 8 tuần", "Tải công việc", "Rủi ro & quyết định", "Kế hoạch tuần"]
assert check["Backlog phân công"].max_row == len(tasks) + 3
assert check["Backlog phân công"]["A4"].value == "P01"
assert check["Tải công việc"]["D4"].value.startswith("=SUMIF")
print(f"Created: {OUTPUT}")
print(f"Tasks: {len(tasks)}")
totals = defaultdict(int)
for t in tasks:
    totals[t[4]] += t[8]
for name, _, _ in members:
    print(f"{name}: {sum(1 for t in tasks if t[4] == name)} tasks, {totals[name]} hours")
