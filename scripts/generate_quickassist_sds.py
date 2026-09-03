from io import BytesIO
from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "QuickAssist_SDS_DacTaThietKePhanMem.docx"

NAVY = "17365D"
BLUE = "2F75B5"
TEAL = "008C95"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"
MID_GRAY = "D9E1F2"
WHITE = "FFFFFF"
RED = "C00000"


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        set_cell_text(hdr.cells[i], h, bold=True, color=WHITE, size=9)
        set_cell_shading(hdr.cells[i], BLUE)
        if widths:
            hdr.cells[i].width = Cm(widths[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
            if widths:
                cells[i].width = Cm(widths[i])
            if ridx % 2 == 1:
                set_cell_shading(cells[i], LIGHT_GRAY)
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_code(doc, text):
    p = doc.add_paragraph(style="Code Block")
    p.add_run(text)
    return p


def add_field(run, instruction):
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("Trang ")
    add_field(paragraph.add_run(), "PAGE")
    paragraph.add_run(" / ")
    add_field(paragraph.add_run(), "NUMPAGES")


def _font(size, bold=False):
    name = "arialbd.ttf" if bold else "arial.ttf"
    try:
        return ImageFont.truetype(name, size)
    except OSError:
        return ImageFont.load_default()


def _center_multiline(draw, box, text, font, fill):
    x1, y1, x2, y2 = box
    bbox = draw.multiline_textbbox((0, 0), text, font=font, align="center", spacing=4)
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.multiline_text(((x1+x2-tw)/2, (y1+y2-th)/2), text, font=font, fill=fill, align="center", spacing=4)


def _arrow(draw, start, end, color="#2F75B5", width=4, both=False):
    import math
    draw.line([start, end], fill=color, width=width)
    def head(tip, tail):
        angle = math.atan2(tip[1]-tail[1], tip[0]-tail[0])
        length = 14
        pts = [tip,
               (tip[0]-length*math.cos(angle-0.55), tip[1]-length*math.sin(angle-0.55)),
               (tip[0]-length*math.cos(angle+0.55), tip[1]-length*math.sin(angle+0.55))]
        draw.polygon(pts, fill=color)
    head(end, start)
    if both:
        head(start, end)


def _polyarrow(draw, points, color="#2F75B5", width=4):
    """Draw a routed one-way arrow so reverse transitions do not overlap."""
    import math
    draw.line(points, fill=color, width=width, joint="curve")
    tip, tail = points[-1], points[-2]
    angle = math.atan2(tip[1]-tail[1], tip[0]-tail[0])
    length = 14
    draw.polygon([
        tip,
        (tip[0]-length*math.cos(angle-0.55), tip[1]-length*math.sin(angle-0.55)),
        (tip[0]-length*math.cos(angle+0.55), tip[1]-length*math.sin(angle+0.55)),
    ], fill=color)


def architecture_image():
    img = Image.new("RGB", (1800, 850), "white")
    draw = ImageDraw.Draw(img)
    draw.text((900, 45), "Kiến trúc logic QuickAssist", font=_font(40, True), fill="#17365D", anchor="mm")
    boxes = [
        ((60, 280, 340, 470), "Người dùng\nChrome / Edge", "#D9EAF7"),
        ((430, 220, 790, 540), "Browser Extension\nUI / Content Script\nService Worker\nIndexedDB", "#C6E0B4"),
        ((900, 220, 1240, 540), "QuickAssist API\nAuth / Notes / Search\nSkim / Quota", "#FFE699"),
        ((1370, 280, 1720, 470), "AI Provider\nChunk / Embed\nRerank", "#F4B183"),
        ((900, 620, 1240, 790), "PostgreSQL\n+ pgvector", "#D9D2E9"),
        ((1370, 620, 1720, 790), "Skim Session\nTTL Store", "#E2F0D9"),
    ]
    for box, label, color in boxes:
        draw.rounded_rectangle(box, radius=22, fill=color, outline="#17365D", width=4)
        _center_multiline(draw, box, label, _font(26, True), "#17365D")
    links = [
        ((340, 375), (430, 375), "Thao tác"),
        ((790, 375), (900, 375), "HTTPS/JSON"),
        ((1240, 375), (1370, 375), "Internal API"),
        ((1070, 540), (1070, 620), "SQL/vector"),
        ((1240, 500), (1480, 620), "Session"),
    ]
    for start, end, label in links:
        _arrow(draw, start, end, both=True)
        mx, my = (start[0]+end[0])//2, (start[1]+end[1])//2
        draw.text((mx, my-20), label, font=_font(18), fill="#595959", anchor="mm")
    bio = BytesIO(); img.save(bio, format="PNG"); bio.seek(0); return bio


def note_state_image():
    # Keep the normal path on the first row and error outcomes on a second row.
    # This prevents branch arrows from visually bypassing SKIPPED/FAILED.
    img = Image.new("RGB", (1800, 740), "white")
    draw = ImageDraw.Draw(img)
    states = [
        (70, 150, "QUEUED", "Chờ 5 giây"),
        (440, 150, "SAVED", "Đã lưu DB"),
        (810, 150, "PROCESSING", "AI đang xử lý"),
        (1180, 150, "READY", "Đã có vector"),
        (650, 485, "SKIPPED", "Không đủ quota"),
        (1110, 485, "FAILED", "Có thể retry"),
    ]
    for x, y, name, detail in states:
        box = (x, y, x+270, y+150)
        draw.rounded_rectangle(box, radius=20, fill="#D9EAF7", outline="#17365D", width=4)
        draw.text((x+135, y+50), name, font=_font(25, True), fill="#17365D", anchor="mm")
        draw.text((x+135, y+102), detail, font=_font(19), fill="#595959", anchor="mm")

    # Main success path: QUEUED → SAVED → PROCESSING → READY.
    for x1, x2 in [(340, 440), (710, 810), (1080, 1180)]:
        _arrow(draw, (x1, 225), (x2, 225))

    # Branches leave PROCESSING and terminate at the appropriate state boxes.
    _arrow(draw, (925, 300), (785, 485), color="#C65911")
    draw.text((740, 392), "Không đủ quota", font=_font(18), fill="#C65911", anchor="mm")
    _arrow(draw, (1035, 300), (1245, 485), color="#C00000")
    draw.text((1290, 392), "Lỗi AI / hạ tầng", font=_font(18), fill="#C00000", anchor="mm")

    # Retry is only a transition from FAILED back to PROCESSING.
    _polyarrow(draw, [(1380, 560), (1530, 560), (1530, 85), (1080, 85), (1080, 150)], color="#70AD47")
    draw.text((1315, 70), "Retry có kiểm soát / idempotency", font=_font(17), fill="#548235", anchor="mm", align="center")

    draw.text((900, 55), "Trạng thái xử lý note", font=_font(34, True), fill="#17365D", anchor="mm")
    bio = BytesIO(); img.save(bio, format="PNG"); bio.seek(0); return bio


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.2)
section.right_margin = Cm(1.8)

# Base styles
styles = doc.styles
styles["Normal"].font.name = "Times New Roman"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
styles["Normal"].font.size = Pt(11)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.15
for style_name, size, color in [("Title", 26, NAVY), ("Heading 1", 16, NAVY), ("Heading 2", 13, BLUE), ("Heading 3", 11.5, TEAL)]:
    style = styles[style_name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
if "Code Block" not in styles:
    code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
else:
    code_style = styles["Code Block"]
code_style.font.name = "Consolas"
code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
code_style.font.size = Pt(8.5)
code_style.paragraph_format.left_indent = Cm(0.5)
code_style.paragraph_format.right_indent = Cm(0.5)
code_style.paragraph_format.space_before = Pt(3)
code_style.paragraph_format.space_after = Pt(6)

# Header/footer
header = section.header.paragraphs[0]
header.text = "QUICKASSIST  |  SOFTWARE DESIGN SPECIFICATION"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    run.font.name = "Arial"; run.font.size = Pt(8); run.font.bold = True; run.font.color.rgb = RGBColor.from_string(BLUE)
footer = section.footer.paragraphs[0]
footer.add_run("QA-SDS-001  •  Nội bộ Nhóm 1                                      ")
add_page_number(footer)
for run in footer.runs:
    run.font.name = "Arial"; run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string("666666")

# Cover
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(78)
r = p.add_run("QUICKASSIST")
r.font.name = "Arial"; r.font.size = Pt(34); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(NAVY)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ĐẶC TẢ THIẾT KẾ PHẦN MỀM")
r.font.name = "Times New Roman"; r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Software Design Specification — SDS")
r.italic = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor.from_string("666666")

doc.add_paragraph("\n")
cover_rows = [
    ("Mã tài liệu", "QA-SDS-001"),
    ("Phiên bản", "1.0 — Bản thiết kế cơ sở"),
    ("Ngày phát hành", "03/09/2026"),
    ("Học phần", "Công nghệ phát triển Phần mềm doanh nghiệp"),
    ("Giảng viên", "TS. Nguyễn Trọng Phúc"),
    ("Nhóm", "Nhóm 1"),
    ("Trạng thái", "DRAFT — Chờ nhóm và giảng viên duyệt"),
]
table = doc.add_table(rows=0, cols=2)
table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
for key, value in cover_rows:
    cells = table.add_row().cells
    set_cell_text(cells[0], key, bold=True, color=NAVY, size=10)
    set_cell_shading(cells[0], LIGHT_BLUE)
    set_cell_text(cells[1], value, size=10)
doc.add_paragraph()
p = doc.add_paragraph("Mai Đức Vinh  •  Trịnh Mạnh Quang  •  Lê Đăng Sơn  •  Trần Thuỳ Dương")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].font.bold = True; p.runs[0].font.color.rgb = RGBColor.from_string(NAVY)
doc.add_page_break()

# Document control
add_heading(doc, "KIỂM SOÁT TÀI LIỆU", 1)
add_heading(doc, "Lịch sử phiên bản", 2)
add_table(doc, ["Phiên bản", "Ngày", "Người soạn", "Nội dung thay đổi", "Trạng thái"], [
    ("0.1", "03/09/2026", "Nhóm 1", "Khởi tạo từ báo cáo ý tưởng và kịch bản chức năng", "Nháp"),
    ("1.0", "03/09/2026", "Nhóm 1", "Hoàn thiện thiết kế cơ sở, API, dữ liệu và kế hoạch kiểm thử", "Chờ duyệt"),
], [2, 3, 4, 8, 3])
add_heading(doc, "Phê duyệt", 2)
add_table(doc, ["Vai trò", "Họ tên", "Ngày", "Chữ ký / xác nhận"], [
    ("Tech Lead", "Mai Đức Vinh", "", ""),
    ("Backend Lead", "Trịnh Mạnh Quang", "", ""),
    ("AI/RAG Lead", "Lê Đăng Sơn", "", ""),
    ("UI/UX & QA Lead", "Trần Thuỳ Dương", "", ""),
    ("Giảng viên", "TS. Nguyễn Trọng Phúc", "", ""),
], [4, 5, 3, 6])
add_heading(doc, "Tài liệu tham chiếu", 2)
add_table(doc, ["Mã", "Tên tài liệu", "Vai trò"], [
    ("REF-01", "Nhóm 1 - Báo cáo ý tưởng đề tài.docx", "Mục tiêu, phạm vi, công nghệ và định vị sản phẩm"),
    ("REF-02", "Kịch bản chức năng QuickAssist_1.xlsx", "Luồng user/UI/server/AI và các nhánh chức năng"),
    ("REF-03", "QuickAssist_KeHoach_PhanCong_4Nguoi.xlsx", "Backlog, ưu tiên, timeline và trách nhiệm"),
], [2.5, 8, 8])
doc.add_page_break()

# TOC
add_heading(doc, "MỤC LỤC", 1)
p = doc.add_paragraph()
add_field(p.add_run(), 'TOC \\o "1-3" \\h \\z \\u')
p2 = doc.add_paragraph("Lưu ý: khi mở trong Microsoft Word, chọn Update Field/Update Table để cập nhật số trang.")
p2.runs[0].italic = True; p2.runs[0].font.color.rgb = RGBColor.from_string("666666")
doc.add_page_break()

# 1 Introduction
add_heading(doc, "1. Giới thiệu", 1)
add_heading(doc, "1.1 Mục đích", 2)
doc.add_paragraph("Tài liệu này mô tả thiết kế phần mềm cho QuickAssist — phần mở rộng trình duyệt hỗ trợ ghi chú nhanh, lưu trữ cá nhân hóa và tra cứu kiến thức bằng tìm kiếm ngữ nghĩa. SDS là cơ sở thống nhất giữa các thành viên khi hiện thực hóa extension, server, AI service, cơ sở dữ liệu và hoạt động kiểm thử.")
add_heading(doc, "1.2 Phạm vi", 2)
add_bullets(doc, [
    "Tài khoản cá nhân, xác thực và quản lý quota sử dụng AI.",
    "Ghi chú nhanh từ đoạn văn bản được chọn, có cửa sổ đếm ngược và hoàn tác.",
    "Thư viện note, phân mục, CRUD, cache IndexedDB và đồng bộ server.",
    "Tra cứu note theo BM25 + vector similarity + reranking.",
    "Tìm văn bản nổi bật trên trang bằng Readability và session tạm thời.",
    "Khả năng cấu hình AI Provider để hỗ trợ dịch vụ trực tuyến hoặc tự triển khai.",
])
doc.add_paragraph("Tóm tắt toàn trang là hạng mục P1 cần hoàn thiện đặc tả trước khi cam kết. Trích xuất lịch trình cá nhân là P2 do còn thiếu quy tắc consent, dữ liệu profile và xử lý thông tin nhạy cảm.")
add_heading(doc, "1.3 Thuật ngữ và viết tắt", 2)
add_table(doc, ["Thuật ngữ", "Diễn giải"], [
    ("SDS", "Software Design Specification — Đặc tả thiết kế phần mềm"),
    ("RAG", "Retrieval-Augmented Generation; trong MVP tập trung retrieval/ranking"),
    ("BM25", "Thuật toán xếp hạng từ khóa dùng trong hybrid search"),
    ("Embedding", "Vector biểu diễn ngữ nghĩa của văn bản"),
    ("Rerank", "Xếp hạng lại tập kết quả ứng viên theo query"),
    ("MV3", "Manifest Version 3 cho Chrome/Edge extension"),
    ("IndexedDB", "Cơ sở dữ liệu cục bộ trong trình duyệt dùng làm cache/offline storage"),
    ("pgvector", "Extension PostgreSQL để lưu và tìm kiếm vector"),
    ("TTL", "Time To Live — thời gian tồn tại của dữ liệu tạm"),
    ("PII", "Personally Identifiable Information — thông tin nhận dạng cá nhân"),
])
add_heading(doc, "1.4 Giả định và ràng buộc", 2)
add_bullets(doc, [
    "Trình duyệt mục tiêu: Chrome/Edge hỗ trợ Manifest V3; Firefox là mục tiêu tương thích sau MVP.",
    "Extension giao tiếp server qua HTTPS; AI endpoint là internal service hoặc provider adapter.",
    "Mỗi dữ liệu note/chunk bắt buộc thuộc đúng một user; không có chia sẻ note trong MVP.",
    "Mọi chỉ tiêu hiệu năng trong SDS là ngưỡng mục tiêu ban đầu và phải hiệu chỉnh bằng benchmark.",
    "Không lưu lâu dài nội dung skim nếu người dùng chưa chọn lưu.",
])

# 2 Overview
add_heading(doc, "2. Tổng quan hệ thống", 1)
add_heading(doc, "2.1 Bối cảnh và mục tiêu", 2)
doc.add_paragraph("QuickAssist nằm giữa utility extension cơ bản và trợ lý AI tích hợp sâu. Sản phẩm giảm thao tác sao chép thủ công, cho phép lưu lại nội dung quan trọng ngay trên trang web, tổ chức theo phân mục và tìm lại bằng ý nghĩa thay vì chỉ bằng từ khóa.")
add_heading(doc, "2.2 Actor", 2)
add_table(doc, ["Actor", "Mô tả", "Quyền chính"], [
    ("Người dùng", "Cá nhân cài extension và có tài khoản", "Tạo/xem/sửa/xóa note của chính mình; tìm kiếm; quản lý phân mục"),
    ("Browser Extension", "Client chạy trên trình duyệt", "Thu thập selection/page text theo thao tác; cache; gọi API"),
    ("QuickAssist Server", "Backend tin cậy", "Xác thực, lưu dữ liệu, áp quota, điều phối AI"),
    ("AI Provider", "Dịch vụ local hoặc remote", "Chunking, embedding, reranking; không tự quyết định quyền truy cập"),
    ("Quản trị vận hành", "Người triển khai hệ thống", "Cấu hình provider, quan sát health/log; không đọc nội dung user theo mặc định"),
])
add_heading(doc, "2.3 Ưu tiên phát hành", 2)
add_table(doc, ["Mức", "Phạm vi", "Tiêu chí"], [
    ("P0 — MVP", "Auth/quota, ghi chú nhanh, thư viện/CRUD/cache, RAG, bảo mật và regression", "Bắt buộc chạy end-to-end và không còn lỗi P0/P1"),
    ("P1", "Skim trang, tóm tắt sau khi bổ sung đặc tả, cấu hình self-host", "Chỉ nhận nếu P0 ổn định"),
    ("P2", "Trích xuất lịch trình cá nhân và tự động xác nhận liên quan", "Cần thiết kế consent/privacy riêng"),
])

# 3 Architecture
add_heading(doc, "3. Kiến trúc phần mềm", 1)
add_heading(doc, "3.1 Nguyên tắc thiết kế", 2)
add_bullets(doc, [
    "Tách rời client, backend và AI qua contract có phiên bản.",
    "Server là điểm thực thi quyền truy cập; AI Provider không nhận định danh user nếu không cần thiết.",
    "Cache-first cho trải nghiệm duyệt note; server vẫn là nguồn dữ liệu chuẩn.",
    "Xử lý AI bất đồng bộ ở nơi phù hợp; trạng thái note phải minh bạch và có thể retry.",
    "Provider-agnostic: Dify, llama.cpp/vLLM hoặc mock được thay bằng cấu hình.",
    "Least privilege, data minimization và deny-by-default.",
])
add_heading(doc, "3.2 Kiến trúc logic", 2)
doc.add_picture(architecture_image(), width=Inches(6.9))
p = doc.paragraphs[-1]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph("Hình 1 — Kiến trúc logic và ranh giới giao tiếp")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].italic = True
add_heading(doc, "3.3 Thành phần và trách nhiệm", 2)
add_table(doc, ["Thành phần", "Trách nhiệm", "Không chịu trách nhiệm"], [
    ("UI/Sidebar/Popup", "Hiển thị, nhận input, trạng thái loading/error/undo", "Không tự xác thực quyền sở hữu dữ liệu"),
    ("Content Script", "Đọc selection, URL, title và nội dung Readability theo quyền", "Không lưu token; không gọi AI trực tiếp"),
    ("Service Worker", "Context menu, shortcut, API client, queue/retry và đồng bộ", "Không lưu dữ liệu dài hạn ngoài local store"),
    ("IndexedDB Store", "Cache note/category, pending queue, sync metadata", "Không thay server làm nguồn chuẩn"),
    ("Auth Module", "Đăng ký/đăng nhập/refresh/logout và xác thực token", "Không xử lý note/search"),
    ("Notes Module", "CRUD, ownership, metadata, trạng thái xử lý", "Không chứa logic mô hình AI"),
    ("Search Module", "Lập search pool, BM25/vector, điều phối rerank", "Không truy cập note user khác"),
    ("Skim Module", "Session tạm có TTL, prepare/search/save selected", "Không lưu page text lâu dài trước consent"),
    ("Quota Module", "Kiểm tra và ghi ledger tiêu hao", "Không quyết định chất lượng kết quả"),
    ("AI Adapter", "Chuẩn hóa chunk/embed/rerank giữa provider", "Không nắm quyền sở hữu dữ liệu"),
    ("PostgreSQL/pgvector", "Dữ liệu chuẩn, transaction và vector index", "Không được client truy cập trực tiếp"),
])
add_heading(doc, "3.4 Công nghệ dự kiến", 2)
add_table(doc, ["Lớp", "Công nghệ", "Lý do"], [
    ("Extension", "TypeScript, Manifest V3, WebExtension APIs", "Phù hợp Chrome/Edge, type safety, đóng gói rõ ràng"),
    ("UI", "TypeScript + framework UI do nhóm chốt", "Component hóa popup/sidebar và trạng thái dùng chung"),
    ("Local storage", "IndexedDB", "Dung lượng lớn hơn localStorage, transaction và index"),
    ("Backend", "Python + framework ASGI", "Tích hợp AI tốt, hỗ trợ async và OpenAPI"),
    ("Database", "PostgreSQL + pgvector", "Quan hệ + vector trong cùng hệ quản trị"),
    ("AI runtime", "Dify hoặc llama.cpp/vLLM qua adapter", "Cho phép cloud/self-host và thay provider"),
    ("Đóng gói", "Docker Compose cho server/DB/AI tùy cấu hình", "Môi trường demo tái lập"),
])
add_heading(doc, "3.5 Kiến trúc triển khai", 2)
add_bullets(doc, [
    "Extension package được cài từ thư mục build hoặc store package; chỉ khai báo permission cần thiết.",
    "Backend chạy sau reverse proxy TLS; chỉ public API được công khai.",
    "PostgreSQL/pgvector đặt trong private network và backup theo lịch.",
    "AI service nội bộ có health check, timeout và giới hạn payload; remote provider đi qua adapter.",
    "Skim session ưu tiên in-memory/Redis-compatible store có TTL; môi trường đơn máy có thể dùng memory store.",
])

# 4 Extension design
add_heading(doc, "4. Thiết kế Browser Extension", 1)
add_heading(doc, "4.1 Cấu trúc module", 2)
add_table(doc, ["Module", "Interface chính", "Dữ liệu"], [
    ("background/service-worker", "registerMenus(), handleCommand(), dispatchRequest(), sync()", "Queue, token reference, sync cursor"),
    ("content/selection", "captureSelection(), expandSelection()", "text, url, title, DOM context tối thiểu"),
    ("content/readability", "extractReadableText()", "clean text, source metadata"),
    ("ui/popup", "showUndo(), showStatus()", "snippet, countdown, action state"),
    ("ui/sidebar", "browseNotes(), search(), skim()", "note DTO, result DTO, pagination"),
    ("storage/indexeddb", "putNote(), queryNotes(), enqueue(), reconcile()", "notes, categories, pending actions, sync metadata"),
    ("api/client", "auth(), notes(), search(), skim()", "versioned JSON DTO"),
])
add_heading(doc, "4.2 Context menu và phím tắt", 2)
add_table(doc, ["Trigger", "Điều kiện", "Hành vi"], [
    ("Ghi chú văn bản này", "Có selection hợp lệ", "Tạo draft local, popup 5 giây, sau đó enqueue request"),
    ("Ctrl+Alt+N", "Có selection", "Tương đương context menu ghi chú nhanh"),
    ("Ctrl+Alt+Z", "Popup undo đang hoạt động", "Hủy timer và loại draft mới nhất khỏi queue"),
    ("Tìm trên trang", "Trang có nội dung đọc được", "Readability extract, tạo skim session và mở UI"),
])
add_heading(doc, "4.3 Thiết kế IndexedDB", 2)
add_table(doc, ["Object store", "Key", "Index", "Mục đích"], [
    ("notes", "noteId", "categoryId, updatedAt, status", "Cache note thuộc user hiện tại"),
    ("categories", "categoryId", "name, updatedAt", "Cache phân mục"),
    ("pendingActions", "actionId", "createdAt, entityId, type", "Queue tạo/sửa/xóa và retry"),
    ("syncState", "resource", "lastSyncAt, cursor", "Theo dõi incremental sync"),
    ("settings", "key", "—", "Tùy chọn AI tự xử lý, notification; không lưu secret thô"),
])
doc.add_paragraph("Quy tắc xung đột mặc định: server-authoritative cho bản ghi đã đồng bộ; optimistic update ở client. Request sửa mang version/updatedAt. Nếu version lệch, server trả 409 và client yêu cầu người dùng tải bản mới hoặc lưu thành bản sao. Xóa dùng tombstone ngắn hạn để đồng bộ giữa thiết bị.")
add_heading(doc, "4.4 Queue, retry và idempotency", 2)
add_bullets(doc, [
    "Mỗi thao tác tạo note sinh actionId/idempotencyKey UUID ở client.",
    "Retry theo exponential backoff có jitter; tối đa 5 lần cho lỗi mạng/5xx.",
    "Không retry tự động cho 400/401/403/409/422; UI phải hướng dẫn hành động.",
    "Server lưu idempotency key trong cửa sổ tối thiểu 24 giờ cho thao tác tạo.",
])

# 5 Backend
add_heading(doc, "5. Thiết kế Backend", 1)
add_heading(doc, "5.1 Phân lớp", 2)
add_table(doc, ["Lớp", "Nhiệm vụ", "Quy tắc phụ thuộc"], [
    ("API/Controller", "Parse/validate request, auth context, status code", "Chỉ gọi application service"),
    ("Application Service", "Use case, transaction boundary, quota orchestration", "Không phụ thuộc framework UI"),
    ("Domain", "Entity, policy ownership/quota/status", "Không phụ thuộc database/provider"),
    ("Repository", "Truy cập PostgreSQL/pgvector", "Implement interface của domain/application"),
    ("Provider Adapter", "Chunk/embed/rerank và timeout/fallback", "Không quyết định authorization"),
])
add_heading(doc, "5.2 Transaction boundary", 2)
add_bullets(doc, [
    "Tạo note và ghi metadata thực hiện trong một transaction ngắn; không giữ transaction DB khi chờ AI.",
    "AI processing cập nhật trạng thái PROCESSING → READY/SKIPPED/FAILED bằng worker hoặc background job.",
    "Chunk và vector của một note được thay thế nguyên tử khi note thay đổi nội dung.",
    "Quota ledger là append-only; cập nhật số dư và ledger trong cùng transaction.",
])
add_heading(doc, "5.3 Quy tắc ownership", 2)
doc.add_paragraph("Mọi repository method truy xuất tài nguyên user phải nhận user_id từ auth context. Không chấp nhận user_id từ JSON để xác định chủ sở hữu. Query note/chunk/category luôn có điều kiện user_id; test bắt buộc chứng minh không truy cập chéo tài khoản.")

# 6 AI
add_heading(doc, "6. Thiết kế AI và RAG", 1)
add_heading(doc, "6.1 Pipeline xử lý note", 2)
add_numbered(doc, [
    "Chuẩn hóa Unicode, loại khoảng trắng thừa và kiểm tra kích thước payload.",
    "Semantic chunking theo câu/đoạn, có overlap cấu hình và giữ source offsets.",
    "Sinh embedding theo model/version đã cấu hình.",
    "Lưu chunk UUID, text, vector, model_version và metadata vào pgvector.",
    "Cập nhật trạng thái note và quota ledger; phát sự kiện trạng thái cho client ở lần sync tiếp theo.",
])
add_heading(doc, "6.2 Pipeline truy vấn", 2)
add_numbered(doc, [
    "Xác thực user và quota, dựng search pool chỉ từ note thuộc user và filter metadata.",
    "Sinh embedding cho query; đồng thời chạy BM25 và vector similarity.",
    "Hợp nhất ứng viên bằng điểm chuẩn hóa/reciprocal-rank fusion theo cấu hình.",
    "Gửi tập ứng viên giới hạn tới reranker; áp timeout và fallback về hybrid score.",
    "Trả top-k chunk kèm noteId, source metadata, score/rank và snippet an toàn.",
])
add_heading(doc, "6.3 Contract provider", 2)
add_table(doc, ["Operation", "Input", "Output", "Timeout mục tiêu", "Fallback"], [
    ("chunk", "text, policy", "chunkId/tempId, text, offsets", "15 giây", "Rule-based paragraph chunking"),
    ("embed", "text[]", "vector[], modelVersion", "10 giây", "Retry/provider dự phòng; không giả vector"),
    ("rerank", "query, candidates[]", "candidateId, score[]", "8 giây", "Giữ hybrid ranking"),
    ("summarize — P1", "clean text, constraints", "summary, citations/sections", "30 giây", "Trả lỗi có thể retry"),
])
add_heading(doc, "6.4 Chỉ số đánh giá", 2)
add_bullets(doc, [
    "Retrieval: Recall@10, MRR@10 hoặc nDCG@10 trên bộ query-note do nhóm gán nhãn.",
    "Latency: p50/p95 theo bước embed, database search, rerank và toàn pipeline.",
    "Robustness: nội dung ngắn/dài, tiếng Việt/Anh, query không có kết quả và provider timeout.",
    "Privacy: payload tới provider không chứa userId/email; log chỉ giữ requestId và số lượng token/ký tự.",
])

# 7 Data
add_heading(doc, "7. Thiết kế dữ liệu", 1)
add_heading(doc, "7.1 Mô hình thực thể", 2)
entities = [
    ("users", "id UUID PK; email CITEXT UNIQUE; password_hash; status; created_at; updated_at", "Tài khoản và trạng thái"),
    ("user_profiles", "user_id PK/FK; display_name; job_title; company; consent_flags JSONB", "Profile tùy chọn; PII tối thiểu"),
    ("categories", "id UUID PK; user_id FK; name; color; created_at; updated_at; version", "Phân mục note"),
    ("notes", "id UUID PK; user_id FK; category_id FK NULL; source_url; source_title; content; status; version; created_at; updated_at; deleted_at", "Dữ liệu note chuẩn"),
    ("note_chunks", "id UUID PK; note_id FK; user_id FK; ordinal; text; embedding VECTOR; model_version; metadata JSONB", "Chunk và vector"),
    ("quota_accounts", "user_id PK; plan; remaining_units; reset_at; updated_at", "Số dư quota hiện hành"),
    ("quota_ledger", "id UUID PK; user_id FK; operation; units; request_id; created_at", "Audit tiêu hao, append-only"),
    ("idempotency_keys", "user_id; key; request_hash; response_ref; expires_at", "Ngăn tạo trùng khi retry"),
    ("skim_sessions", "session_id UUID; user_id; tab_ref_hash; status; expires_at", "Metadata session; text tạm ở TTL store"),
]
add_table(doc, ["Bảng", "Trường chính", "Mục đích"], entities, [3.5, 10, 5])
add_heading(doc, "7.2 Index và ràng buộc", 2)
add_bullets(doc, [
    "UNIQUE(users.email), UNIQUE(user_id, categories.name) khi chưa xóa.",
    "INDEX notes(user_id, updated_at DESC), notes(user_id, category_id, updated_at DESC).",
    "INDEX note_chunks(user_id, note_id); vector index HNSW/IVFFlat được chọn sau benchmark kích thước dữ liệu.",
    "CHECK status thuộc tập trạng thái hợp lệ; version tăng khi sửa.",
    "Foreign key cascade cho note_chunks khi xóa cứng note; note ưu tiên soft delete trong thời gian đồng bộ.",
])
add_heading(doc, "7.3 Vòng đời dữ liệu", 2)
add_table(doc, ["Dữ liệu", "Tạo", "Lưu", "Xóa/retention"], [
    ("Note", "User chọn lưu", "Đến khi user xóa", "Soft delete để sync; purge theo policy cấu hình"),
    ("Chunk/vector", "Sau khi AI xử lý", "Theo note", "Xóa/rebuild cùng note hoặc đổi model"),
    ("Skim page text", "Khi prepare", "TTL ngắn, đề xuất 30 phút", "Tự hết hạn; chỉ đoạn được chọn mới thành note"),
    ("Quota ledger", "Mỗi operation AI", "Theo audit policy", "Không chứa nội dung văn bản"),
    ("Log", "Trong vận hành", "Đề xuất 14–30 ngày", "Redact token, PII và nội dung"),
])

# 8 API
add_heading(doc, "8. Thiết kế API", 1)
add_heading(doc, "8.1 Quy ước chung", 2)
add_bullets(doc, [
    "Base path: /api/v1; JSON UTF-8; thời gian ISO-8601 UTC; ID dùng UUID.",
    "Authorization: Bearer access token; refresh token lưu an toàn theo phương án triển khai.",
    "Request tạo/sửa nhận X-Request-ID; tạo note nhận Idempotency-Key.",
    "Pagination dùng cursor; filter phải được whitelist và giới hạn.",
    "Lỗi theo một schema thống nhất; không trả stack trace cho client.",
])
add_code(doc, '''{
  "error": {
    "code": "QUOTA_EXCEEDED",
    "message": "Không đủ quota để xử lý AI.",
    "requestId": "uuid",
    "details": {"operation": "embedding"}
  }
}''')
add_heading(doc, "8.2 Public API", 2)
api_rows = [
    ("POST", "/auth/register", "—", "email, password", "201 User; 409 email tồn tại"),
    ("POST", "/auth/login", "—", "email, password", "200 tokens/user; 401"),
    ("POST", "/auth/refresh", "Refresh", "refresh token", "200 access token; 401"),
    ("POST", "/auth/logout", "User", "refresh/session ref", "204"),
    ("POST", "/notes/quick", "User", "content, source, categoryId", "202 Note(status=SAVED/PROCESSING)"),
    ("GET", "/notes", "User", "cursor, updatedAfter, categoryId, status", "200 items + nextCursor"),
    ("GET", "/notes/{noteId}", "Owner", "—", "200 Note; 404"),
    ("PUT", "/notes/{noteId}", "Owner", "content/category/version", "200 Note; 409 version conflict"),
    ("DELETE", "/notes/{noteId}", "Owner", "—", "204"),
    ("POST", "/notes/bulk-delete", "Owner", "noteIds/category selector + confirm", "202/204"),
    ("POST", "/search/ranking", "User+quota", "query, categoryId?, topK?", "200 ranked chunks"),
    ("POST", "/skim/sessions", "User+quota", "source, cleanText", "202 sessionId/status"),
    ("GET", "/skim/sessions/{id}", "Owner", "—", "200 status; 404/410"),
    ("POST", "/skim/sessions/{id}/search", "Owner+quota", "query, topK", "200 ranked chunks"),
    ("POST", "/notes/from-skim", "Owner", "sessionId, selectedChunkIds, categoryId", "201 notes"),
    ("POST", "/summaries — P1", "User+quota", "source, cleanText, save", "202 summary; chỉ sau khi G01 duyệt"),
]
add_table(doc, ["Method", "Path", "Quyền", "Input chính", "Kết quả"], api_rows, [2, 5, 3.5, 6, 5])
doc.add_paragraph("Quyết định chuẩn hóa: dùng /notes/quick cho tạo note từ selection và /notes/from-skim cho lưu đoạn từ skim. Backend không tiếp tục dùng tên chung /notes hoặc /notes/ragged cho hai luồng khác nhau nếu chưa có operation discriminator rõ ràng.")
add_heading(doc, "8.3 Internal AI API", 2)
add_table(doc, ["Method/Path", "Input", "Output", "Bảo vệ"], [
    ("POST /internal/ai/chunk", "text, policy", "chunks[]", "Private network/service credential"),
    ("POST /internal/ai/embed", "texts[]", "vectors[], modelVersion", "Payload/rate limit; timeout"),
    ("POST /internal/ai/rerank", "query, candidates[]", "ranked ids/scores", "Candidate cap; timeout"),
    ("GET /internal/ai/health", "—", "provider/model/status", "Không public chi tiết secret"),
])
add_heading(doc, "8.4 Mã trạng thái", 2)
add_table(doc, ["HTTP", "Ý nghĩa sử dụng"], [
    ("200/201", "Thành công đồng bộ / đã tạo"),
    ("202", "Đã nhận và sẽ xử lý bất đồng bộ"),
    ("204", "Thành công không có body"),
    ("400/422", "Request sai định dạng / validation domain"),
    ("401/403", "Chưa xác thực / không có quyền"),
    ("404", "Không tồn tại hoặc cố ý che tài nguyên không thuộc user"),
    ("409", "Version/idempotency conflict"),
    ("410", "Skim session đã hết hạn"),
    ("429", "Rate limit hoặc quota vượt giới hạn"),
    ("502/503/504", "AI provider/service unavailable/timeout"),
])

# 9 workflows
add_heading(doc, "9. Thiết kế luồng chức năng", 1)
add_heading(doc, "9.1 Ghi chú nhanh", 2)
add_numbered(doc, [
    "User bôi đen văn bản và chọn menu hoặc Ctrl+Alt+N.",
    "Content Script lấy selection, URL, title; UI hiển thị snippet và đếm ngược 5 giây.",
    "Nếu user hoàn tác, timer và pending action bị hủy; không gửi request.",
    "Nếu hết thời gian, Service Worker ghi pending action rồi POST /api/v1/notes/quick với idempotency key.",
    "Server xác thực, kiểm tra ownership/category và lưu note. Nếu bật AI và đủ quota, trạng thái chuyển PROCESSING.",
    "AI chunk/embed; server lưu vector, ghi quota ledger và chuyển READY. Hết quota chuyển SKIPPED, lỗi chuyển FAILED.",
    "Client đồng bộ trạng thái vào IndexedDB và thông báo phù hợp.",
])
doc.add_picture(note_state_image(), width=Inches(6.9))
p = doc.paragraphs[-1]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph("Hình 2 — Trạng thái xử lý note")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].italic = True
add_heading(doc, "9.2 Duyệt và CRUD note", 2)
add_bullets(doc, [
    "UI query IndexedDB trước để hiển thị nhanh; sau đó incremental sync bằng updatedAfter/cursor.",
    "Sửa note dùng optimistic update và version; server trả 409 nếu xung đột.",
    "Khi nội dung thay đổi, trạng thái vector chuyển PROCESSING và chunk cũ chỉ được thay khi bộ mới hoàn chỉnh.",
    "Xóa số lượng lớn hoặc cả phân mục phải có xác nhận rõ số lượng; request server hoàn tất trước khi purge cache.",
])
add_heading(doc, "9.3 Tra cứu note RAG-like", 2)
add_numbered(doc, [
    "User nhập query trong toàn bộ thư viện hoặc một phân mục.",
    "Client POST /search/ranking; server xác thực và dựng pool theo user/category.",
    "Server embed query, chạy BM25 + vector, hợp nhất và rerank.",
    "Server ghi quota ledger và trả top-k chunk kèm metadata nguồn.",
    "UI ưu tiên render dữ liệu cache nếu note nguồn đã có; nếu thiếu thì đồng bộ note cần thiết.",
])
add_heading(doc, "9.4 Tìm văn bản nổi bật trên trang", 2)
add_numbered(doc, [
    "Extension chạy Readability và tạo sessionRef UUID theo tab + user.",
    "POST /skim/sessions gửi cleanText; server kiểm tra quota rồi xử lý vào TTL store.",
    "Trong lúc chuẩn bị, UI cho nhập query nhưng hiển thị trạng thái 'Đang hiểu nội dung trang web'.",
    "Khi READY, client gửi query; pipeline hybrid + rerank trả các đoạn nổi bật.",
    "User preview và chọn đoạn; POST /notes/from-skim chuyển các đoạn được chọn thành note lâu dài.",
])
add_heading(doc, "9.5 Tóm tắt nội dung — thiết kế dự kiến P1", 2)
doc.add_paragraph("Chức năng này có trong báo cáo ý tưởng nhưng chưa có kịch bản chi tiết. Thiết kế dưới đây là placeholder và không được xem là contract cuối cùng cho đến khi nhóm duyệt G01.")
add_table(doc, ["Vấn đề cần chốt", "Đề xuất"], [
    ("Đầu vào", "Clean text từ Readability; giới hạn kích thước và cảnh báo trang nhạy cảm"),
    ("Đầu ra", "Summary có cấu trúc, source URL/title, model version và thời điểm tạo"),
    ("Lưu trữ", "Mặc định chỉ lưu khi user chọn; không tự đưa vào RAG"),
    ("Quota", "Tính theo operation và kích thước input; hiển thị lỗi rõ"),
    ("Lịch trình", "Tắt trong MVP; chỉ bật sau consent/profile specification"),
])

# 10 Security
add_heading(doc, "10. Bảo mật và quyền riêng tư", 1)
add_heading(doc, "10.1 Threat model tóm tắt", 2)
add_table(doc, ["Mối đe dọa", "Kiểm soát thiết kế"], [
    ("Rò token từ extension", "Không log token; storage phù hợp; CSP; không nhúng secret provider"),
    ("Truy cập note user khác", "Auth context + repository filter user_id + test IDOR"),
    ("XSS từ nội dung trang/note", "Render text an toàn; sanitize nếu cho phép rich text; cấm innerHTML tùy ý"),
    ("Prompt injection trong page text", "AI chỉ chunk/embed/rerank; nội dung trang là untrusted; không cấp tool/action"),
    ("SSRF/URL abuse", "Server không fetch URL do client gửi trong MVP; URL chỉ là metadata"),
    ("Payload quá lớn", "Giới hạn ký tự/body, rate limit, timeout và streaming nếu cần"),
    ("Data leakage tới provider", "Data minimization, redact PII, cấu hình local, DPA/consent khi dùng remote"),
    ("Xóa hàng loạt ngoài ý muốn", "Xác nhận số lượng, ownership, audit và soft delete"),
])
add_heading(doc, "10.2 Xác thực và mật khẩu", 2)
add_bullets(doc, [
    "Mật khẩu hash bằng Argon2id hoặc bcrypt với tham số theo môi trường; không tự thiết kế thuật toán hash.",
    "Access token ngắn hạn; refresh token có rotation/revocation. Logout vô hiệu session/refresh token.",
    "TLS bắt buộc ngoài localhost. CORS chỉ cho origin extension và web quản trị được cấu hình.",
    "Rate limit cho login, notes, skim và search; quota không thay thế rate limit.",
])
add_heading(doc, "10.3 Privacy", 2)
add_bullets(doc, [
    "Chỉ thu thập nội dung sau hành động rõ ràng của user; hiển thị trang/đoạn sắp gửi khi khả thi.",
    "Không gửi email, userId, company/job title sang AI nếu không cần cho use case.",
    "Có chức năng xóa note và dữ liệu vector liên quan; ghi rõ retention cho skim/log/backup.",
    "Lịch trình cá nhân yêu cầu consent riêng, có thể rút lại và có chế độ xác nhận thủ công.",
])

# 11 NFR
add_heading(doc, "11. Yêu cầu phi chức năng", 1)
add_table(doc, ["Mã", "Nhóm", "Yêu cầu thiết kế / mục tiêu ban đầu", "Cách đo"], [
    ("NFR-01", "Hiệu năng", "Popup phản hồi thao tác < 200 ms; API CRUD p95 < 500 ms, không tính mạng", "Browser/API benchmark"),
    ("NFR-02", "AI latency", "Quick-note processing p95 < 15 s; RAG p95 < 8 s ở cấu hình demo", "Telemetry theo stage"),
    ("NFR-03", "Tin cậy", "Không tạo note trùng khi retry; action queue khôi phục sau restart", "Fault injection/E2E"),
    ("NFR-04", "Offline", "Duyệt cache và enqueue note khi mất mạng; sync sau reconnect", "E2E offline"),
    ("NFR-05", "Bảo mật", "100% endpoint dữ liệu có auth + ownership test", "Security integration suite"),
    ("NFR-06", "Quy mô demo", "10.000 note/user, 100.000 chunk/user không làm sai phân quyền", "Load test dữ liệu sinh"),
    ("NFR-07", "Khả dụng", "Loading/empty/error/retry rõ; thao tác chính dùng được bằng bàn phím", "Usability/accessibility review"),
    ("NFR-08", "Tương thích", "Chrome/Edge phiên bản hỗ trợ MV3 tại thời điểm phát hành", "Compatibility matrix"),
    ("NFR-09", "Quan sát", "Mọi request có requestId; log không chứa token/nội dung note", "Log inspection"),
    ("NFR-10", "Bảo trì", "API có OpenAPI; module có owner; migration có rollback/forward plan", "Release checklist"),
])

# 12 errors/observability
add_heading(doc, "12. Xử lý lỗi và quan sát hệ thống", 1)
add_heading(doc, "12.1 Phân loại lỗi", 2)
add_table(doc, ["Loại", "Ví dụ", "Hành vi client", "Hành vi server"], [
    ("Validation", "Selection rỗng, payload quá dài", "Hiển thị ngay; không retry", "400/422 + field details"),
    ("Auth", "Token hết hạn", "Refresh một lần; thất bại thì login", "401; không lộ tài nguyên"),
    ("Quota", "Không đủ embedding/rerank", "Giải thích note vẫn được lưu hoặc search bị chặn", "429/operation code; ledger nhất quán"),
    ("Network/5xx", "Mất mạng/provider lỗi", "Queue/retry hoặc nút thử lại", "Timeout, circuit breaker/fallback"),
    ("Conflict", "Sửa note version cũ", "Yêu cầu reload/lưu bản sao", "409 + currentVersion"),
    ("Session", "Skim hết TTL", "Cho chuẩn bị lại", "410 Gone"),
])
add_heading(doc, "12.2 Logging và metric", 2)
add_bullets(doc, [
    "Structured log: timestamp, level, requestId, operation, duration, status, provider/model version; không ghi raw content.",
    "Metrics: request count/error/latency; queue depth; AI latency/failure; DB/vector latency; quota rejection.",
    "Health: liveness tách readiness; readiness kiểm tra DB và cấu hình provider nhưng tránh gọi model nặng mỗi lần.",
])

# 13 deployment
add_heading(doc, "13. Cấu hình và triển khai", 1)
add_heading(doc, "13.1 Biến cấu hình", 2)
add_table(doc, ["Nhóm", "Ví dụ", "Quy tắc"], [
    ("Database", "DATABASE_URL, VECTOR_DIMENSION", "Secret qua environment/secret store; dimension khớp model"),
    ("Auth", "TOKEN_TTL, REFRESH_TTL, PASSWORD_HASH_PARAMS", "Không dùng giá trị demo trong production"),
    ("AI", "AI_PROVIDER, MODEL_NAME, AI_BASE_URL, TIMEOUT", "Provider key chỉ ở server"),
    ("Quota", "DEFAULT_PLAN, OPERATION_COSTS, RESET_POLICY", "Version hóa chính sách"),
    ("Skim", "SESSION_TTL, MAX_PAGE_CHARS", "Giới hạn để bảo vệ bộ nhớ và chi phí"),
    ("CORS/Extension", "ALLOWED_EXTENSION_IDS, ALLOWED_ORIGINS", "Whitelist theo môi trường"),
])
add_heading(doc, "13.2 Quy trình triển khai", 2)
add_numbered(doc, [
    "Chuẩn bị secret/config và kiểm tra version PostgreSQL/pgvector.",
    "Chạy migration theo phiên bản; xác minh backup/restore ở môi trường demo.",
    "Khởi động DB → backend → AI adapter/provider; kiểm tra readiness.",
    "Build extension bằng API base URL của môi trường; kiểm tra manifest permission.",
    "Chạy smoke test auth, quick note, CRUD, RAG và skim nếu bật.",
    "Ghi release version, migration version, model version và known issues.",
])

# 14 Testing
add_heading(doc, "14. Chiến lược kiểm thử", 1)
add_table(doc, ["Cấp", "Phạm vi", "Owner chính", "Điều kiện pass"], [
    ("Unit", "Parser, domain policy, scoring, storage adapters", "Mỗi module", "Nhánh quan trọng và edge case pass"),
    ("Contract", "Extension–API và Server–AI DTO", "Vinh/Quang/Sơn", "OpenAPI/schema tương thích"),
    ("Integration", "DB transaction, ownership, quota, vector search", "Quang/Sơn", "Môi trường test tái lập"),
    ("E2E", "5 luồng và các nhánh lỗi từ Excel", "Dương", "Không lỗi P0/P1 mở"),
    ("Security", "IDOR, token, XSS, rate/payload limit", "Vinh/Quang", "Checklist và automated test pass"),
    ("AI evaluation", "Retrieval quality, latency, provider fallback", "Sơn", "Đạt ngưỡng đã chốt"),
    ("Usability", "Keyboard, loading/error/undo, demo flow", "Dương", "Người thử hoàn thành task cốt lõi"),
])
add_heading(doc, "14.1 Test scenario bắt buộc", 2)
add_bullets(doc, [
    "Ghi note thành công; undo trong 5 giây; hết quota nhưng note vẫn lưu; mất mạng và retry không tạo trùng.",
    "Browse từ cache; cache miss sync server; sửa conflict; xóa một note và xóa hàng loạt.",
    "RAG toàn thư viện và theo category; query không kết quả; provider rerank timeout; kiểm tra chéo user.",
    "Skim chưa READY, trang dài, session hết hạn, hết quota và lưu nhiều đoạn được chọn.",
    "Token hết hạn, logout, password sai, payload quá lớn, source text chứa HTML/script.",
])

# 15 Traceability
add_heading(doc, "15. Ma trận truy vết", 1)
trace_rows = [
    ("FR-01", "Đăng ký/đăng nhập", "Auth Module, users/quota_accounts", "/auth/*", "AUTH-E2E"),
    ("FR-02", "Ghi chú nhanh", "Selection/Popup/Queue/Notes/AI", "/notes/quick", "QN-E2E"),
    ("FR-03", "Hoàn tác 5 giây", "Popup + pendingActions", "Không gọi API khi undo", "QN-UNDO"),
    ("FR-04", "Duyệt note/cache", "Sidebar/IndexedDB/Notes", "GET /notes", "CACHE-SYNC"),
    ("FR-05", "Sửa/xóa note", "Sidebar/Notes/Versioning", "PUT/DELETE /notes/{id}", "CRUD-E2E"),
    ("FR-06", "Tra cứu RAG", "Search/AI Adapter/pgvector", "POST /search/ranking", "RAG-E2E/EVAL"),
    ("FR-07", "Tìm trên trang", "Readability/Skim/Search", "/skim/sessions/*", "SKIM-E2E"),
    ("FR-08", "Lưu đoạn từ skim", "Skim/Notes", "POST /notes/from-skim", "SKIM-SAVE"),
    ("FR-09", "Quota", "Quota middleware/ledger", "Áp dụng notes/search/skim", "QUOTA-INT"),
    ("FR-10", "Tóm tắt", "Summary placeholder", "POST /summaries — P1", "Chờ G01"),
]
add_table(doc, ["Mã", "Yêu cầu", "Thành phần", "API", "Test"], trace_rows, [2, 4.5, 6, 5.5, 3])

# 16 Decisions/open issues
add_heading(doc, "16. Quyết định thiết kế và vấn đề mở", 1)
add_heading(doc, "16.1 Quyết định đã đề xuất", 2)
add_table(doc, ["ADR", "Quyết định", "Lý do"], [
    ("ADR-001", "Server-authoritative + IndexedDB cache-first", "UI nhanh nhưng quyền và dữ liệu chuẩn vẫn ở server"),
    ("ADR-002", "Chuẩn hóa /notes/quick và /notes/from-skim", "Loại bỏ mâu thuẫn endpoint trong kịch bản"),
    ("ADR-003", "Provider adapter cho Dify/local/mock", "Giảm khóa nền tảng và tăng khả năng demo"),
    ("ADR-004", "AI processing tách transaction lưu note", "Tránh giữ transaction dài và vẫn lưu note khi AI lỗi"),
    ("ADR-005", "Skim text chỉ lưu tạm TTL", "Giảm rủi ro privacy và dung lượng"),
    ("ADR-006", "Hybrid retrieval có rerank fallback", "Giữ kết quả khi provider rerank timeout"),
])
add_heading(doc, "16.2 Vấn đề cần nhóm phê duyệt", 2)
open_items = [
    ("OI-01", "Framework UI extension cụ thể", "Vinh/Dương", "Cuối tuần 1"),
    ("OI-02", "Framework ASGI/ORM và thư viện migration", "Quang", "Cuối tuần 1"),
    ("OI-03", "Dify hay service tự quản cho demo", "Sơn", "Cuối tuần 1"),
    ("OI-04", "Ngưỡng quota và đơn vị tính", "Quang/Sơn", "Tuần 2"),
    ("OI-05", "Quy tắc sync conflict và retention soft delete", "Vinh/Quang", "Tuần 2"),
    ("OI-06", "Scenario/acceptance cho tóm tắt", "Dương", "Tuần 2"),
    ("OI-07", "Consent/profile cho lịch trình cá nhân", "Cả nhóm", "Trước khi đưa vào P2"),
    ("OI-08", "Ngưỡng relevance/latency chính thức", "Sơn/Dương", "Sau benchmark tuần 5"),
]
add_table(doc, ["Mã", "Nội dung", "Owner", "Hạn chốt"], open_items, [2.5, 10, 4.5, 4])

# Appendix
add_heading(doc, "Phụ lục A — Cấu trúc payload tham khảo", 1)
add_heading(doc, "A.1 Tạo note nhanh", 2)
add_code(doc, '''POST /api/v1/notes/quick
Authorization: Bearer <access-token>
Idempotency-Key: <uuid>

{
  "content": "Đoạn văn bản người dùng đã chọn",
  "source": {
    "url": "https://example.com/article",
    "title": "Article title"
  },
  "categoryId": "uuid-or-null",
  "autoProcess": true
}''')
add_heading(doc, "A.2 Kết quả tìm kiếm", 2)
add_code(doc, '''{
  "requestId": "uuid",
  "items": [
    {
      "chunkId": "uuid",
      "noteId": "uuid",
      "text": "Đoạn nội dung liên quan",
      "rank": 1,
      "score": 0.91,
      "source": {"url": "https://...", "title": "..."},
      "categoryId": "uuid-or-null"
    }
  ],
  "quota": {"remainingUnits": 120}
}''')
add_heading(doc, "Phụ lục B — Checklist review SDS", 1)
add_bullets(doc, [
    "Phạm vi P0/P1/P2 đã được cả nhóm xác nhận.",
    "API và data model không còn mâu thuẫn với kịch bản chức năng.",
    "Mỗi use case có owner, acceptance criteria và test tương ứng.",
    "Quyền riêng tư cho page text, profile và AI provider đã được chấp thuận.",
    "Các chỉ tiêu NFR đã được benchmark hoặc ghi rõ là mục tiêu tạm thời.",
    "Các open item OI-01…OI-08 có quyết định/ADR trước khi triển khai phụ thuộc.",
])

# Final document settings
for paragraph in doc.paragraphs:
    paragraph.paragraph_format.widow_control = True
    for run in paragraph.runs:
        if run.font.name is None:
            run.font.name = "Times New Roman"

doc.core_properties.title = "QuickAssist — Đặc tả thiết kế phần mềm (SDS)"
doc.core_properties.subject = "Software Design Specification cho QuickAssist"
doc.core_properties.author = "Nhóm 1"
doc.core_properties.keywords = "QuickAssist, SDS, Browser Extension, RAG, pgvector, Manifest V3"
doc.core_properties.comments = "Sinh từ báo cáo ý tưởng, kịch bản chức năng và kế hoạch phân công của Nhóm 1."
doc.save(OUTPUT)

# Reopen for structural validation.
check = Document(OUTPUT)
headings = [p.text for p in check.paragraphs if p.style.name.startswith("Heading")]
assert len(check.tables) >= 20
assert "8. Thiết kế API" in headings
assert "15. Ma trận truy vết" in headings
assert len(check.inline_shapes) >= 2
print(f"Created: {OUTPUT}")
print(f"Paragraphs: {len(check.paragraphs)}")
print(f"Tables: {len(check.tables)}")
print(f"Headings: {len(headings)}")
print(f"Figures: {len(check.inline_shapes)}")
